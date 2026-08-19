"""
Generate a driller whose revenue block carries an offset line — the shape
that asks a tile to show a net figure: Banking NII and Wealth Fees stand
positive under Revenue, Other Revenue stands negative beside them, and the
cost lines stand on their own under Costs. Figures are random; the shape is
the point.
"""
import random
from openpyxl import Workbook

random.seed(41)

DIM_HEADERS = [
    "MICA", "MICA_Level_3", "MICA_Level_2", "MICA_Level_1",
    "Product_code", "Product_Level_3", "Product_Level_2", "Product_Level_1",
    "Segment_code", "CG_Level_2", "CG_Level_1",
    "Function_code", "Function_Level_2", "Function_Level_1",
    "Entity code",
]

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

PERIODS = (
    [(f"{m}-25", "Actuals") for m in MONTHS] +
    [(f"{q}Q25 Actual", "Actuals") for q in range(1, 5)] +
    [(f"{m}-26", "Actuals") for m in MONTHS[:6]] +
    [(f"{m}-26", "Forecast") for m in MONTHS[6:]] +
    [(f"{q}Q26 Actual", "Actuals") for q in range(1, 3)] +
    [("JUN YTD-25", "Actuals"), ("JUN YTD-26", "Actuals"),
     ("JUN YTD-26 Target", "Target"),
     ("FY-25", "Actuals"), ("FY-26 Forecast", "Forecast"), ("FY-26 Target", "Target")]
)

# code, level 3, level 2, level 1, sign, scale
MICA_ROWS = [
    ("MP10101010000", "NII - Interest Income",  "Banking NII",   "Revenue", +1, 60),
    ("MP10102010000", "NFI - Fee Income",       "Wealth Fees",   "Revenue", +1, 34),
    ("MP10104010000", "Revenue offsets",        "Other Revenue", "Revenue", -1, 14),
    ("MP20101010000", "Staff Costs",            "Direct Costs",  "Costs",   -1, 26),
    ("MP20201010000", "Recharges",              "Indirect Costs", "Costs",  -1, 11),
]

PRODUCTS = [
    ("PR04020000", "Mortgages (Real Estate Secured)", "Loans"),
    ("PR04030000", "Credit Cards",                    "Cards"),
    ("PR05001000", "Current Accounts Total",          "Deposits"),
    ("PR12100000", "Insurance Distribution",          "Insurance"),
]

SEGMENT = ("CG1010000", "Retail Banking & Wealth Management", "IWPB")
FUNCTIONS = [
    ("RTN00008", "Premier and Wealth", "Intl Wealth & Premier Banking"),
    ("RTN00012", "Personal Banking",   "Intl Wealth & Premier Banking"),
]
ENTITY = "4040_1MGT"


def val(base, sign):
    return round(sign * abs(random.gauss(base, base * 0.2)), 1)


wb = Workbook()
ws = wb.active
ws.title = "Sheet1"
ws.append([""] * len(DIM_HEADERS) + [b for _, b in PERIODS])
ws.append(DIM_HEADERS + [h for h, _ in PERIODS])

for mica, l3, l2, l1, sign, scale in MICA_ROWS:
    for pcode, p3, p2 in PRODUCTS:
        for fcode, f2, f1 in FUNCTIONS:
            base = random.uniform(scale * 0.6, scale * 1.4)
            drift = random.uniform(0.9, 1.18)
            py = [val(base, sign) for _ in range(12)]
            pyq = [round(sum(py[i * 3:i * 3 + 3]), 1) for i in range(4)]
            cy = [val(base * drift, sign) for _ in range(6)]
            fc = [val(base * drift, sign) for _ in range(6)]
            cyq = [round(sum(cy[0:3]), 1), round(sum(cy[3:6]), 1)]
            ytd_py = round(sum(py[:6]), 1)
            ytd_cy = round(sum(cy), 1)
            ytd_tgt = round(ytd_cy * random.uniform(0.94, 1.06), 1)
            fy25 = round(sum(py), 1)
            fy26fc = round(sum(cy) + sum(fc), 1)
            fy26tgt = round(fy26fc * random.uniform(0.93, 1.07), 1)
            dims = [mica, l3, l2, l1, pcode, p3, p2, "Total Product",
                    SEGMENT[0], SEGMENT[1], SEGMENT[2], fcode, f2, f1, ENTITY]
            ws.append(dims + py + pyq + cy + fc + cyq +
                      [ytd_py, ytd_cy, ytd_tgt, fy25, fy26fc, fy26tgt])

out = "IWPB_SG_Driller_offsets.xlsx"
wb.save(out)
print(f"wrote {out}: {ws.max_row-2} rows x {ws.max_column} cols")
