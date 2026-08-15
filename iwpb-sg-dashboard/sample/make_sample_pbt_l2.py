"""
Generate a driller shaped like the real IWPB product extract in the field:
MICA Level 1 is a single 'PBT' for every row — the P&L result — and Level 2
carries the seven distinct lines the tiles should read: Revenue, Total Direct
Cost, Total Indirect Costs (incl InterCo), Expected credit losses, Share of
Profit in Associates & JV's, and the two balance memo lines Customer Deposits
and Loan and Advances. Level 3/4 split Revenue into NII / Net Fee Income the
way the real file does. Figures are random; the shape is the point.
"""
import random
from openpyxl import Workbook

random.seed(7)

DIM_HEADERS = [
    "MICA", "MICA_Level_4", "MICA_Level_3", "MICA_Level_2", "MICA_Level_1",
    "Product_code", "Product_Level_3", "Product_Level_2", "Product_Level_1",
    "Segment_code", "CG_Level_2", "CG_Level_1",
    "Entity code",
]

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

PERIODS = (
    [(f"{m}-25", "Actuals") for m in MONTHS] +
    [(f"{q}Q25 Actual", "Actuals") for q in range(1, 5)] +
    [(f"{m}-26", "Actuals") for m in MONTHS[:6]] +
    [(f"{m}-26", "Forecast") for m in MONTHS[6:]] +
    [(f"{q}Q26 Actual", "Actuals") for q in range(1, 3)] +
    [("JUN YTD-25", "Actuals"), ("JUN YTD-26", "Actuals"),
     ("JUN YTD-26 Target", "Target"),
     ("FY-25", "Actuals"), ("FY-26 Forecast", "Forecast"), ("FY-26 Target", "Target")]
)

# code, level 4, level 3, level 2 — level 1 is always PBT; sign, scale, kind
MICA_ROWS = [
    ("MP10101010000", "NII - Interest Income",  "NII - Net Interest Income", "Revenue", +1, 55, "flow"),
    ("MP10101020000", "NII - Interest Expense", "NII - Net Interest Income", "Revenue", -1, 22, "flow"),
    ("MP10201010000", "NFI - Fee Income",       "Net Fee Income",            "Revenue", +1, 30, "flow"),
    ("MP20101010000", "Staff Costs",            "Direct Costs",              "Total Direct Cost", -1, 24, "flow"),
    ("MP20201010000", "Recharges",              "Indirect Costs",            "Total Indirect Costs (incl InterCo)", -1, 12, "flow"),
    ("MP30101010000", "ECL Charge",             "ECL",                       "Expected credit losses", -1, 6, "flow"),
    ("MP40101010000", "JV Share",               "Associates",                "Share of Profit in Associates & JV's", +1, 2, "flow"),
    ("MP90101010000", "Deposits balance",       "Deposits",                  "Customer Deposits", +1, 900, "bal"),
    ("MP90201010000", "Loans balance",          "Loans",                     "Loan and Advances", +1, 800, "bal"),
]

PRODUCTS = [
    ("PR04004001", "Wealth Lending",                  "Loans"),
    ("PR04020000", "Mortgages (Real Estate Secured)", "Loans"),
    ("PR05001000", "Current Accounts Total",          "Deposits"),
    ("PR12200000", "Insurance Distribution",          "Insurance"),
]

SEGMENT = ("CG1010000", "Retail Banking & Wealth Management", "IWPB")
ENTITY = "4040_1MGT"


wb = Workbook()
ws = wb.active
ws.title = "IWPB India"
ws.append([""] * len(DIM_HEADERS) + [b for _, b in PERIODS])
ws.append(DIM_HEADERS + [h for h, _ in PERIODS])

for mica, l4, l3, l2, sign, scale, kind in MICA_ROWS:
    for pcode, p3, p2 in PRODUCTS:
        base = random.uniform(scale * 0.7, scale * 1.3)
        drift = random.uniform(0.92, 1.15)
        if kind == "bal":
            # a balance column carries the closing balance, month after month
            def bal(seed):
                return [round(seed * random.uniform(0.97, 1.03), 1) for _ in range(12)]
            py = bal(base)
            cy = bal(base * drift)[:6]
            fc = bal(base * drift)[:6]
            pyq = [py[2], py[5], py[8], py[11]]
            cyq = [cy[2], cy[5]]
            ytd_py, ytd_cy = py[5], cy[5]
            ytd_tgt = round(cy[5] * random.uniform(0.95, 1.05), 1)
            fy25 = py[11]
            fy26fc = fc[5]
            fy26tgt = round(fc[5] * random.uniform(0.95, 1.05), 1)
        else:
            py = [round(sign * abs(random.gauss(base, base * .2)), 1) for _ in range(12)]
            cy = [round(sign * abs(random.gauss(base * drift, base * .2)), 1) for _ in range(6)]
            fc = [round(sign * abs(random.gauss(base * drift, base * .2)), 1) for _ in range(6)]
            pyq = [round(sum(py[i*3:i*3+3]), 1) for i in range(4)]
            cyq = [round(sum(cy[0:3]), 1), round(sum(cy[3:6]), 1)]
            ytd_py = round(sum(py[:6]), 1)
            ytd_cy = round(sum(cy), 1)
            ytd_tgt = round(ytd_cy * random.uniform(0.94, 1.06), 1)
            fy25 = round(sum(py), 1)
            fy26fc = round(sum(cy) + sum(fc), 1)
            fy26tgt = round(fy26fc * random.uniform(0.93, 1.07), 1)
        dims = [mica, l4, l3, l2, "PBT", pcode, p3, p2, "Total Product",
                SEGMENT[0], SEGMENT[1], SEGMENT[2], ENTITY]
        ws.append(dims + py + pyq + cy + fc + cyq +
                  [ytd_py, ytd_cy, ytd_tgt, fy25, fy26fc, fy26tgt])

out = "IWPB_PBT_L2_Driller.xlsx"
wb.save(out)
print(f"wrote {out}: {ws.max_row-2} rows x {ws.max_column} cols")
