# ============================================================
# REFERENCE DATA / HEADCOUNT / APPIAN CONSOLIDATION
# FULL CONSOLIDATED SCRIPT
#
# Fixes in this version:
#   1. FRC reconciliation now uses ALL Appian rows (previously rows
#      without a Position ID were dropped before the reconciliation).
#   2. apply_data_types only strips a trailing ".0" and only converts
#      a column to numbers/dates when no value would be lost.
#   3. Rows with a blank Position Number are kept (previously all but
#      one were silently dropped by drop_duplicates).
#   4. FRC reconciliation no longer uppercases FRC Code in the
#      exported Reference file.
#   5. Pre/Post comparison is text-based, so 1234 vs "1234" is not
#      reported as a change.
#   6. HIRING_FLG can optionally be reset each run (RESET_HIRING_FLG).
#   7. Removed positions can optionally be dropped from the Reference
#      output (DROP_REMOVED_POSITIONS).
#   8. ID columns are formatted as "0" instead of "#,##0".
#   9. Columns are converted to object dtype before cell-by-cell
#      updates (avoids pandas 2.x warnings / pandas 3 errors).
#  10. Job Summary exclusions can be maintained on the Config sheet.
#
# Comments_RD rules (applied in this order):
#   a. Vacant last month, occupied this month      -> "Onboarded"
#   b. STATUS_APPROVAL = AWAITING COO-BE-LEAD APPROVAL
#      or AWAITING APPROVAL FROM OPCO              -> "AWAITING FRC APPROVAL"
#   c. STATUS_APPROVAL = OPCO APPROVED and
#      Comments_RD blank                           -> "RAISE IJP"
#
# Input files: leave the C:\YOUR_PATH placeholders in section 1 and a
# Windows "Open" box asks for each file when the script runs.
#
# Leaver rule (applies after every other rule, no exceptions):
#   Vacant position (STATUS = VACANT or no employee) -> Leaver blank;
#   a vacant position can never have LEAVER = YES
# Vacant last month = Employee Name Vacant/blank or Status = VACANT in
# the PM file; filled = real Employee Name in HC for the SAME Position
# Number (see the "Filled-Changed check" sheet).
# Filled position rule (vacant last month and occupied this month, OR
# a different employee in the position this month):
#   Vacant last month, occupied this month         -> Committed Offer,
#   Candidate Name, Joiner flag, Joiner Month, Comments_RD and Leaver
#   blanked, Status flag set to PHYSICAL
# Vacated position rule:
#   Occupied last month, vacant this month         -> Employee Name
#   "Vacant", Status flag set to VACANT, Leaver cleared
# HC names: built from Employee First + Last Name; blank or
# "Unspecified" in HC means the position is vacant.
# Position attributes (GCB, Position Title, Work Location, ...):
#   HC "Position Global Career Band" is read as the Reference's GCB
#   column (other spellings likewise), so new positions get a GCB;
#   a BLANK in HC / Appian never wipes the value already in the
#   Reference (kept values listed on "Blank values kept").
# Committed offer rule:
#   Comments_RD = OFFERED                          -> Committed Offer = YES
#   (all Committed Offer = YES positions listed on "Committed Offers")
# Sub Function Head rule:
#   Position in HC, Sub Function Head Status BLANK in PM and
#   Completed in Appian this month                 -> HIRING_FLG, LEAVER
#                                                     and JOINER = YES
#                                                     (LEAVER blank if vacant);
#   and if Status = PHYSICAL                       -> HIRING AGAINST
#                                                     PHYSICAL_RD = YES
# DEMISE rule (Hyderabad):
#   Location like Hyderabad, Sub Function Head Status = Completed now,
#   position in HC, HIRING_FLG not already DEMISE in PM (instead of
#   the hiring outcome above)             -> HIRING_FLG = DEMISE,
#     NEW_POSITION_FLAG = DEMISE, LEAVER = YES (blank if vacant),
#     Comments_RD = "Raise New Position" (or "DEMISE" when the position
#     is vacant), HIRING AGAINST PHYSICAL_RD blank. Every Hyderabad
#     position is listed on the "DEMISE check" sheet with the outcome.
#   (DEMISE_ONLY_NEWLY_APPROVED = True: only where blank in the PM file)
# Hiring flag rule (HIRING_FLG and HIRING AGAINST PHYSICAL_RD):
#   Vacant last month with flag = YES, filled
#   this month                                     -> flag removed
# New position rule:
#   Position Number not in last month's Reference  -> NEW_POSITION_FLAG = YES
#   (existing flag values are never overwritten)
#   Status flag set to PHYSICAL (employee in HC) or VACANT (no
#   employee in HC); same for an existing position whose Status flag
#   is blank in the PM file
# LWD rule:
#   Current month Employee ID = Employee ID (MR_Leavers_CM)
#                                                  -> LWD = Termination Date
#   (only Employee Class = Employee; exited employees / Vacant
#    positions are not updated)
#
# Outputs (all in the OUTPUT folder): Reference, Exception Report,
# MR Hiring Update, FRC Reconciliation and Run_Log_<timestamp>.txt
# ============================================================

import os
import re
import sys

import numpy as np
import pandas as pd

from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# 1. INPUT FILES / OUTPUT FOLDER / OPTIONS
# ============================================================

reference_file = r"C:\YOUR_PATH\Reference_Data_Hierarchy.xlsx"
hc_file = r"C:\YOUR_PATH\HC_Current_Month.xlsx"
appian_file = r"C:\YOUR_PATH\Appian.xlsx"
# Leavers file - its Termination Date updates LWD (section 9G).
# Optional: if this file is not found, the LWD update is skipped.
leavers_file = r"C:\YOUR_PATH\MR_Leavers_CM.xlsx"
# Outputs go to an OUTPUT folder next to this script.
# To use a different folder, replace this line with e.g.
#   output_folder = r"C:\YOUR_PATH\OUTPUT"
output_folder = os.path.join(os.path.dirname(os.path.abspath(__file__)), "OUTPUT")

# True  -> HIRING_FLG is cleared and re-derived every run
# False -> existing YES values are kept (original behaviour)
RESET_HIRING_FLG = False

# True  -> positions no longer in HC are removed from the Reference output
# False -> they stay in the Reference output (original behaviour);
#          they are always listed in the Exception Report
DROP_REMOVED_POSITIONS = False

# True  -> a blank GCB / Position Title / Work Location / ... in the HC
#          or Appian file does NOT wipe the value already in the
#          Reference (kept values are listed on "Blank values kept")
# False -> blanks in the source files overwrite the Reference
KEEP_VALUE_WHEN_SOURCE_BLANK = True   # <-- change to False to turn OFF

# True  -> the DEMISE outcome (Hyderabad) only for positions whose Sub
#          Function Head Status was blank in the PM file and is Completed
#          now (same trigger as the hiring outcome)
# False -> every Hyderabad position in HC with Sub Function Head Status =
#          Completed now, unless HIRING_FLG is already DEMISE in the PM file
DEMISE_ONLY_NEWLY_APPROVED = False    # <-- change to True to restrict


# ------------------------------------------------------------
# FILE PICKER
#
# If a path above is still the C:\YOUR_PATH placeholder, a Windows
# "Open" box pops up to choose that file - no need to type paths.
# ------------------------------------------------------------

def pick_file(title):
    try:
        import tkinter as tk
        from tkinter import filedialog

        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)  # show the box in front of Jupyter
        path = filedialog.askopenfilename(
            title=title,
            filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")],
        )
        root.destroy()
    except Exception as err:
        raise RuntimeError(
            f"Could not open a file picker ({err}). "
            "Type the file paths in section 1 instead."
        )
    if not path:
        raise RuntimeError(f"No file chosen for: {title}")
    return path


if "YOUR_PATH" in reference_file:
    reference_file = pick_file("Select the REFERENCE DATA HIERARCHY file")
if "YOUR_PATH" in hc_file:
    hc_file = pick_file("Select the HC CURRENT MONTH file")
if "YOUR_PATH" in appian_file:
    appian_file = pick_file("Select the APPIAN file")

os.makedirs(output_folder, exist_ok=True)

timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")

output_reference_file = os.path.join(
    output_folder, f"Reference_Data_Hierarchy_CM_{timestamp}.xlsx"
)
exception_file = os.path.join(
    output_folder, f"Exception_Report_{timestamp}.xlsx"
)
hiring_update_file = os.path.join(
    output_folder, f"MR_Hiring_Update_{timestamp}.xlsx"
)
frc_reconciliation_file = os.path.join(
    output_folder, f"FRC_Reconciliation_Report_{timestamp}.xlsx"
)
run_log_file = os.path.join(
    output_folder, f"Run_Log_{timestamp}.txt"
)


# ------------------------------------------------------------
# RUN LOG
#
# Everything printed (counts, warnings, errors) is shown on screen
# AND written to Run_Log_<timestamp>.txt in the output folder.
# ------------------------------------------------------------

class TeeToLog:
    def __init__(self, stream, log_path):
        self.stream = stream
        self.log_path = log_path

    def write(self, text):
        self.stream.write(text)
        # Open/append/close each time so the log is complete
        # even if the run stops with an error
        with open(self.log_path, "a", encoding="utf-8") as f:
            f.write(text)

    def flush(self):
        self.stream.flush()

    def __getattr__(self, name):
        return getattr(self.stream, name)


def original_stream(stream):
    # Unwrap a log set up by an earlier run (e.g. re-running a notebook cell)
    while hasattr(stream, "log_path"):
        stream = stream.stream
    return stream


sys.stdout = TeeToLog(original_stream(sys.stdout), run_log_file)
sys.stderr = TeeToLog(original_stream(sys.stderr), run_log_file)

print(f"Run started: {datetime.now():%d-%b-%Y %H:%M:%S}")
print(f"Reference file : {reference_file}")
print(f"HC file        : {hc_file}")
print(f"Appian file    : {appian_file}")


# ============================================================
# 2. COLUMN DEFINITIONS
# ============================================================

KEY_COL = "Position Number"
EMPLOYEE_NAME_COL = "Employee Name"
OLD_EMP_RD_COL = "OLD_EMP_RD"
COMMENTS_RD_COL = "Comments_RD"
ONBOARDED_COMMENT = "Onboarded"
NEW_POSITION_FLAG_COL = "NEW_POSITION_FLAG"
# Other spellings of the same field (case, spaces, "_" and "-" are
# ignored when matching), so the existing column is always reused.
NEW_POSITION_FLAG_CANDIDATES = [
    NEW_POSITION_FLAG_COL, "NEW_POSITION_FLG", "NEW_POS_FLAG", "NEW_POS_FLG",
    "NEW_POSITION", "NEW_POSITION_YN",
]
# Status column (e.g. STATUS, STATUS_FLAG) - spelling-tolerant match
STATUS_CANDIDATES = ["STATUS", "STATUS_FLAG", "STATUS_FLG", "POSITION STATUS", "STATUS_RD"]
FILLED_STATUS_VALUE = "PHYSICAL"   # status given to a position once it is filled
VACANT_STATUS_VALUE = "VACANT"     # status given to a position that became vacant
# Names / IDs that mean "no employee" (the HC file uses "Unspecified")
NO_EMPLOYEE_VALUES = ["vacant", "unspecified"]
# Leaver column: exact names first, else any column starting with "LEAVER"
LEAVER_CANDIDATES = ["LEAVER", "LEAVERS", "LEAVER_RD", "LEAVER_FLG", "LEAVER FLAG", "LEAVER_FLAG_RD"]
OFFERED_COMMENT = "OFFERED"          # Comments_RD value that means an offer is committed
COMMITTED_OFFER_CANDIDATES = ["COMMITTED OFFER", "COMMITTED_OFFER_FLAG", "COMMITTED OFFER FLG"]
CANDIDATE_NAME_CANDIDATES = ["CANDIDATE NAME", "CANDIDATE", "CANDIDATE_NAME_RD"]
JOINER_MONTH_CANDIDATES = ["JOINER MONTH", "JOINERS MONTH", "JOINING MONTH", "JOINER_MONTH_RD"]
MR_LEAD_CANDIDATES = ["MR_LEAD", "MR LEAD", "MR_LEAD_RD", "MR LEAD NAME", "MR_LEAD_NAME"]
# Sub Function Head approved AND Work Location in DEMISE_LOCATIONS ->
# DEMISE instead of hiring: HIRING_FLG and NEW_POSITION_FLAG = DEMISE,
# LEAVER = YES (blank if vacant), Comments_RD = DEMISE_COMMENT,
# HIRING AGAINST PHYSICAL_RD blank
# Location "like" Hyderabad: any location column (Work Location, Location,
# City, Site ...) containing one of these, any case, e.g. "Hyderabad -
# Hitec City", "IN-HYDERABAD". Names of 3 letters or fewer ("HYD") must be
# a whole word, so "Hyde Park" does not match.
DEMISE_LOCATIONS = ["Hyderabad", "HYD"]
DEMISE_VALUE = "DEMISE"
DEMISE_COMMENT = "Raise New Position"   # Comments_RD when the position is occupied
DEMISE_VACANT_COMMENT = "DEMISE"        # Comments_RD when the position is vacant
GCB_CANDIDATES = [
    "GCB", "Global Career Band", "GCB Level", "Position Global Career Band",
    "Position Career Band", "Career Band",
]
WORK_LOCATION_CANDIDATES = ["Work Location", "Location", "Work Location Name", "Position Location"]
POSITION_TITLE_CANDIDATES = ["Position Title", "Position Name"]
# Same field, different column name in the HC file (e.g. HC "Position
# Global Career Band" = Reference "GCB"): the HC column is read under
# the Reference's name, so new positions get the value and existing
# positions are compared.
HC_COLUMN_ALIASES = {
    "GCB": GCB_CANDIDATES,
    "Work Location": WORK_LOCATION_CANDIDATES,
    "Position Title": POSITION_TITLE_CANDIDATES,
}
# Position attributes: a BLANK in the HC / Appian file does not
# overwrite the value already in the Reference (see
# KEEP_VALUE_WHEN_SOURCE_BLANK). Employee fields (Employee ID, Employee
# Name, ...) are deliberately NOT listed, so a vacated position is
# still cleared.  "*" = any column starting with the name.
KEEP_WHEN_SOURCE_BLANK = GCB_CANDIDATES + WORK_LOCATION_CANDIDATES + POSITION_TITLE_CANDIDATES + [
    "Job Summary", "FRC Code", "Job Code", "Job Category", "Workforce Plan ID",
    "Functional Manager*",
]
JOB_SUMMARY_COL = "Job Summary"
# Matched ignoring case, extra spaces, and "&" vs "and"
EXPECTED_JOB_SUMMARY = "Financial Insight & Advisory Support Specialist"

FRC_COL = "FRC Code"
POSITION_ID_COL = "Position ID"

NUMERIC_COLUMNS = [
    "Position Number",
    "Position ID",
    "Existing Position ID",
    "Employee ID",
    "GCB",
    "GCB Level",
    "Global Career Band",
    "Position Global Career Band",
    "Position Career Band",
    "Functional Manager Position Level Position Number",
    "Functional Manager Job Level Employee ID",
]

DATE_COLUMNS = [
    "EXIT DATE",
    "Exit Date",
    "JOINER MONTH",
    "Joiners Month",
    "JOINER_MONTH_RD",
    "LWD",
    "LWD_Pre",
    "LWD_Post",
    "Termination Date",
    "Terminate Date",
]

DEFAULT_JOB_SUMMARY_EXCLUSIONS = [
    "Operations Leadership",
    "Business Finance - Modelling",
    "Bus Finance Asst - Modelling",
    "Risk Analytics and Modelling",
]


# ============================================================
# 3. HELPER FUNCTIONS
# ============================================================

def clean_columns(df):
    df = df.copy()
    df.columns = [str(c).replace("\xa0", " ").strip() for c in df.columns]
    return df


def clean_compare_value(value):
    if pd.isna(value):
        return ""
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return value.strip()
    return value


def values_differ(old_val, new_val):
    # Compare as text so 1234 (number) and "1234" (text) are equal
    return str(clean_compare_value(old_val)) != str(clean_compare_value(new_val))


def clean_position_id_series(series):
    def clean_one(value):
        if pd.isna(value):
            return np.nan
        value = str(value).strip()
        if value == "" or value.lower() in ["nan", "none", "<na>"]:
            return np.nan
        # Remove trailing .0 from Excel numeric IDs
        if re.fullmatch(r"\d+\.0", value):
            value = value[:-2]
        # Numeric IDs: ignore spaces and leading zeros ("0051866881" = "51866881")
        digits = value.replace(" ", "")
        if digits.isdigit():
            value = digits.lstrip("0") or "0"
        return value

    return series.apply(clean_one).astype(object)


def clean_frc_series(series):
    return (
        series.fillna("")
        .astype(str)
        .str.replace("\xa0", " ", regex=False)
        .str.strip()
        .str.upper()
    )


def clean_text_series(series, lower=False):
    s = series.fillna("").astype(str).str.strip()
    return s.str.lower() if lower else s.str.upper()


def normalise_text(value):
    # Lower case, "&" read as "and", repeated / non-breaking spaces collapsed
    if pd.isna(value):
        return ""
    text = str(value).replace("\xa0", " ").lower().replace("&", " and ")
    return " ".join(text.split())


def is_blank_series(series):
    s = series.astype(str).str.strip().str.lower()
    return series.isna() | s.isin(["", "nan", "none", "nat", "<na>"])


def find_column(df, name):
    # Case-insensitive column match, e.g. "comments_rd" -> "Comments_RD"
    for col in df.columns:
        if str(col).strip().lower() == name.strip().lower():
            return col
    return None


def find_column_loose(df, candidates):
    # Like find_column, but also ignores spaces, "_" and "-",
    # e.g. "New Position Flag" matches "NEW_POSITION_FLAG"
    def key(name):
        return re.sub(r"[^a-z0-9]", "", str(name).lower())
    for name in candidates:
        for col in df.columns:
            if key(col) == key(name):
                return col
    return None


def find_column_prefix(df, candidates, prefix):
    # Exact (loose) match first, else the first column starting with prefix
    col = find_column_loose(df, candidates)
    if col is not None:
        return col
    key = re.sub(r"[^a-z0-9]", "", prefix.lower())
    for c in df.columns:
        if re.sub(r"[^a-z0-9]", "", str(c).lower()).startswith(key):
            return c
    return None


def find_functional_manager_col(df):
    # Name column first (e.g. "Functional Manager Job Level Employee Name"),
    # else any Functional Manager column that is not an ID / number
    for c in df.columns:
        if str(c).lower().startswith("functional manager") and "name" in str(c).lower():
            return c
    for c in df.columns:
        if (str(c).lower().startswith("functional manager")
                and not re.search(r"\bid\b|number", str(c).lower())):
            return c
    return None


def keep_when_blank(col):
    # Position-attribute column whose Reference value is kept when the
    # HC / Appian value is blank (see KEEP_WHEN_SOURCE_BLANK)
    key = re.sub(r"[^a-z0-9]", "", str(col).lower())
    for name in KEEP_WHEN_SOURCE_BLANK:
        k = re.sub(r"[^a-z0-9]", "", name.lower())
        if key == k or (name.endswith("*") and key.startswith(k)):
            return True
    return False

def is_demise_location(value):
    # "Hyderabad - Hitec City", "HYDERABAD", "IN-HYD" ... -> True
    text = normalise_text(value)
    tokens = re.split(r"[^a-z0-9]+", text)
    for name in DEMISE_LOCATIONS:
        n = normalise_text(name)
        if len(n) <= 3:
            if n in tokens:
                return True
        elif n in text:
            return True
    return False


def location_columns(df):
    # Every column that looks like a location (Work Location, Location,
    # City, Site, Office ...)
    return [c for c in df.columns if re.search(r"location|city|site|office", str(c).lower())]


def is_vacant(value):
    # "Vacant" or "Unspecified" (HC wording) both mean no employee
    return str(clean_compare_value(value)).strip().lower() in NO_EMPLOYEE_VALUES


def split_blank_keys(df, key_col):
    # Rows with no key cannot be indexed/de-duplicated safely,
    # so they are set aside and added back afterwards.
    has_key = df[key_col].notna()
    return df[has_key].copy(), df[~has_key].copy()


def drop_duplicate_keys(df, key_col, label):
    dupes = df[key_col].duplicated(keep="first")
    if dupes.any():
        print(
            f"WARNING: {dupes.sum()} duplicate {key_col} rows in {label} "
            "- only the first occurrence is used."
        )
    return df[~dupes].copy()


def apply_data_types(df):
    df = df.copy()

    for col in NUMERIC_COLUMNS:
        if col not in df.columns:
            continue

        cleaned = (
            df[col]
            .astype(str)
            .str.strip()
            .str.replace(r"\.0$", "", regex=True)
        )
        numeric = pd.to_numeric(cleaned, errors="coerce")
        blank = is_blank_series(df[col])

        # Only convert when nothing would be lost
        # (e.g. an alphanumeric ID keeps the column as text)
        lost = numeric.isna() & ~blank
        if numeric.notna().any() and not lost.any():
            df[col] = numeric.astype("Int64")
        elif lost.any():
            # Mixed column: numeric IDs become numbers,
            # alphanumeric IDs are kept exactly as they are
            numeric_ok = numeric.notna() & (numeric % 1 == 0)
            df[col] = df[col].astype(object)
            df.loc[numeric_ok, col] = numeric[numeric_ok].astype("int64")
            df.loc[blank, col] = None
            print(
                f"Note: '{col}' has {lost.sum()} non-numeric value(s) "
                "- those are kept as text."
            )

    for col in DATE_COLUMNS:
        if col not in df.columns:
            continue

        parsed = [to_date(v) for v in df[col]]
        lost = pd.Series([v is None for v in parsed], index=df.index)
        if not lost.any():
            df[col] = pd.to_datetime(pd.Series(parsed, index=df.index, dtype=object))
        else:
            print(
                f"Note: '{col}' kept as text - {lost.sum()} "
                "value(s) could not be read as dates."
            )

    return df


def to_date(value):
    # One cell -> date. 0, blanks and placeholder dates (01-Jan-1970,
    # 00-Jan-1900) become blank; Excel serial numbers are converted.
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return pd.NaT
    if isinstance(value, str):
        value = value.strip()
        if value.lower() in ["", "nan", "none", "nat", "<na>"]:
            return pd.NaT
        if re.fullmatch(r"\d+(\.0+)?", value):
            value = float(value)
    if isinstance(value, (int, float, np.integer, np.floating)) and not isinstance(value, bool):
        if value <= 0:
            return pd.NaT
        date = pd.Timestamp("1899-12-30") + pd.Timedelta(days=float(value))
    else:
        try:
            date = pd.to_datetime(value)
        except (TypeError, ValueError):
            return None  # not a date
    if pd.isna(date) or date.year <= 1970:
        return pd.NaT
    return date


HIGHLIGHT_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")  # light green


def highlight_cells(file_path, sheet_name, column_name, keys, key_col=None):
    # Light-green fill on <column_name> for rows whose key is in <keys>
    if not os.path.exists(file_path) or not keys:
        return 0
    key_col = key_col or KEY_COL
    wb = load_workbook(file_path)
    if sheet_name not in wb.sheetnames:
        return 0
    ws = wb[sheet_name]
    headers = {str(ws.cell(1, c).value).strip(): c for c in range(1, ws.max_column + 1)}
    if column_name not in headers or key_col not in headers:
        return 0
    keys = {str(k) for k in keys}
    count = 0
    for r in range(2, ws.max_row + 1):
        key = str(clean_compare_value(ws.cell(r, headers[key_col]).value))
        if key in keys:
            ws.cell(r, headers[column_name]).fill = HIGHLIGHT_FILL
            count += 1
    wb.save(file_path)
    return count


def format_excel(file_path):
    if not os.path.exists(file_path):
        return

    wb = load_workbook(file_path)

    header_fill = PatternFill(fill_type="solid", fgColor="DB0011")
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        if ws.max_row < 1:
            continue

        # Header formatting
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.freeze_panes = "A2"

        if ws.max_column > 0:
            ws.auto_filter.ref = ws.dimensions

        # Auto width
        for col_idx in range(1, ws.max_column + 1):
            max_length = 0
            for row_idx in range(1, min(ws.max_row, 5000) + 1):
                value = ws.cell(row=row_idx, column=col_idx).value
                if value is not None:
                    max_length = max(max_length, len(str(value)))
            width = min(max(max_length + 2, 10), 50)
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # Numeric/date formatting
        header_map = {
            ws.cell(1, col).value: col for col in range(1, ws.max_column + 1)
        }

        for header in NUMERIC_COLUMNS:
            if header in header_map:
                col_idx = header_map[header]
                for row_idx in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    if isinstance(cell.value, (int, float)):
                        # IDs: no thousands separator
                        cell.number_format = "0"

        for header in DATE_COLUMNS:
            if header in header_map:
                col_idx = header_map[header]
                for row_idx in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    if isinstance(cell.value, (datetime, pd.Timestamp)):
                        cell.number_format = "dd-mmm-yyyy"

    wb.save(file_path)


# ============================================================
# 4. CONFIG HELPERS
# ============================================================

def get_config_list(config_df, candidate_columns):
    for col in candidate_columns:
        if col in config_df.columns:
            values = []
            for value in config_df[col]:
                if pd.notna(value):
                    value = str(value).strip()
                    if value:
                        values.append(value)
            if values:
                return values
    return []


def get_export_column_order(config_df):
    return get_config_list(
        config_df, ["Fields", "Field", "Column", "Column Name"]
    )


def apply_column_sequence(df, requested_order):
    if not requested_order:
        return df
    first = [c for c in requested_order if c in df.columns]
    remaining = [c for c in df.columns if c not in first]
    return df[first + remaining]


def get_appian_config(appian_config_df):
    config = clean_columns(appian_config_df)
    field_col = config.columns[0]

    fields = []
    primary_key = None

    for _, row in config.iterrows():
        field = row.get(field_col)
        if pd.isna(field):
            continue
        field = str(field).strip()
        if not field:
            continue

        fields.append(field)

        # Second column contains Primary_Key marker
        if len(config.columns) > 1:
            marker = row.get(config.columns[1])
            if pd.notna(marker) and str(marker).strip().lower() == "primary_key":
                primary_key = field

    if primary_key is None:
        primary_key = "Position ID"

    return primary_key, fields


def get_job_summary_exclusions(config_df):
    # Add a "Job Summary Exclusions" column on the Config sheet to
    # maintain this list without editing the script.
    exclusions = get_config_list(
        config_df,
        ["Job Summary Exclusions", "Job Summary Exclusion", "JOB_SUMMARY_EXCLUSIONS"],
    ) or DEFAULT_JOB_SUMMARY_EXCLUSIONS

    return {normalise_text(x) for x in exclusions}


def get_hiring_update_columns(config_df):
    # If not available, all columns are used later.
    return get_config_list(
        config_df, ["Hiring Update", "MR Hiring Update", "Hiring_Update"]
    )


# ============================================================
# 5. READ INPUT FILES
# ============================================================

reference_xls = pd.ExcelFile(reference_file)
reference_sheet = reference_xls.sheet_names[0]

for candidate in ["Reference_Data", "Reference Data", "Reference_Data_PM"]:
    if candidate in reference_xls.sheet_names:
        reference_sheet = candidate
        break

ref_df = pd.read_excel(reference_file, sheet_name=reference_sheet)
config_df = pd.read_excel(reference_file, sheet_name="Config")
appian_config_df = pd.read_excel(reference_file, sheet_name="APPIAN_CONFIG")
new_migrations_df = pd.read_excel(reference_file, sheet_name="New_Migrations")

hc_df = pd.read_excel(hc_file)
appian_df = pd.read_excel(appian_file)

ref_df = clean_columns(ref_df)
config_df = clean_columns(config_df)
appian_config_df = clean_columns(appian_config_df)
new_migrations_df = clean_columns(new_migrations_df)
hc_df = clean_columns(hc_df)

# ------------------------------------------------------------
# HC EMPLOYEE NAME
# The HC report has "Employee First Name" + "Employee Last Name"
# instead of "Employee Name". Build "Employee Name" from them so the
# Reference is updated, and use "Vacant" when the HC shows no employee
# (blank / "Unspecified").
# ------------------------------------------------------------
hc_name_col = find_column_loose(hc_df, [EMPLOYEE_NAME_COL, "Employee Full Name", "Full Name"])
hc_first_col = find_column_loose(hc_df, ["Employee First Name", "First Name"])
hc_last_col = find_column_loose(hc_df, ["Employee Last Name", "Last Name"])


def hc_name_part(series):
    s = series.fillna("").astype(str).str.strip()
    return s.where(~s.str.lower().isin(NO_EMPLOYEE_VALUES + ["nan"]), "")


if hc_name_col is not None:
    names = hc_name_part(hc_df[hc_name_col])
elif hc_first_col is not None or hc_last_col is not None:
    first = hc_name_part(hc_df[hc_first_col]) if hc_first_col else ""
    last = hc_name_part(hc_df[hc_last_col]) if hc_last_col else ""
    names = (first + " " + last).str.strip()
    print(f"HC Employee Name built from '{hc_first_col}' + '{hc_last_col}'")
else:
    names = None
    print("WARNING: no Employee Name / First Name / Last Name column in HC.")

if names is not None:
    if hc_name_col is not None and hc_name_col != EMPLOYEE_NAME_COL:
        hc_df = hc_df.drop(columns=[hc_name_col])
    hc_df[EMPLOYEE_NAME_COL] = names.where(names != "", "Vacant")

# ------------------------------------------------------------
# HC COLUMN NAMES -> REFERENCE COLUMN NAMES
# e.g. HC "Position Global Career Band" is the Reference's "GCB".
# Without this, new positions would get a blank GCB (the HC column
# name does not match) and existing positions would never be compared.
# ------------------------------------------------------------
for _label, _aliases in HC_COLUMN_ALIASES.items():
    _ref_col = find_column_loose(ref_df, _aliases)
    if _ref_col is None or _ref_col in hc_df.columns:
        continue  # no such column in the Reference / HC already uses the same name
    _hc_col = find_column_loose(hc_df, [_ref_col] + _aliases)
    if _hc_col is not None and _hc_col not in ref_df.columns:
        hc_df = hc_df.rename(columns={_hc_col: _ref_col})
        print(f"HC column '{_hc_col}' read as Reference column '{_ref_col}'")

appian_df = clean_columns(appian_df)


# ============================================================
# 6. CLEAN PRIMARY KEYS
# ============================================================

ref_df[KEY_COL] = clean_position_id_series(ref_df[KEY_COL])
hc_df[KEY_COL] = clean_position_id_series(hc_df[KEY_COL])

hc_blank_keys = hc_df[KEY_COL].isna().sum()
if hc_blank_keys:
    print(
        f"WARNING: {hc_blank_keys} HC row(s) have a blank {KEY_COL} "
        "and cannot be matched."
    )


# ============================================================
# 7. IDENTIFY REMOVED / COMMON / NEW POSITIONS
# ============================================================

ref_keys = set(ref_df[KEY_COL].dropna())
hc_keys = set(hc_df[KEY_COL].dropna())

removed_keys = ref_keys - hc_keys
new_keys = hc_keys - ref_keys
common_keys = ref_keys & hc_keys

positions_removed_df = ref_df[ref_df[KEY_COL].isin(removed_keys)].copy()


# ============================================================
# 8. NEW POSITIONS ADDED FROM CM
# ============================================================

# NEW_POSITION_FLAG: "YES" for positions that did not exist last month.
# The EXISTING flag column in the Reference is used (no new column is
# ever created). Flags already present (on existing rows, or a value
# already on the new row) are kept, never overwritten.
updated_ref_df = ref_df.copy()
new_flag_col = find_column_loose(updated_ref_df, NEW_POSITION_FLAG_CANDIDATES)

new_positions_df = hc_df[hc_df[KEY_COL].isin(new_keys)].copy()
new_positions_df = drop_duplicate_keys(new_positions_df, KEY_COL, "HC new positions")

if new_flag_col is None:
    print(f"WARNING: no '{NEW_POSITION_FLAG_COL}' column in the Reference - "
          "new positions not flagged (no column added).")
else:
    print(f"New position flag column: '{new_flag_col}'")
    # If HC also carries the flag (any spelling), line it up with the Reference column
    hc_flag_col = find_column_loose(new_positions_df, NEW_POSITION_FLAG_CANDIDATES)
    if hc_flag_col is not None and hc_flag_col != new_flag_col:
        new_positions_df = new_positions_df.drop(columns=[new_flag_col], errors="ignore")
        new_positions_df = new_positions_df.rename(columns={hc_flag_col: new_flag_col})
    if new_flag_col not in new_positions_df.columns:
        new_positions_df[new_flag_col] = ""
    new_positions_df[new_flag_col] = new_positions_df[new_flag_col].astype(object)
    flag_blank = is_blank_series(new_positions_df[new_flag_col])
    new_positions_df.loc[flag_blank, new_flag_col] = "YES"
    print(f"New positions this month: {len(new_positions_df)} "
          f"({int(flag_blank.sum())} flagged YES, {int((~flag_blank).sum())} kept existing flag)")

for col in updated_ref_df.columns:
    if col not in new_positions_df.columns:
        new_positions_df[col] = ""

new_positions_df = new_positions_df[updated_ref_df.columns]

updated_ref_df = pd.concat([updated_ref_df, new_positions_df], ignore_index=True)

# Object dtype so cells can take any value during the updates below
updated_ref_df = updated_ref_df.astype(object)

# Rows with a blank Position Number are set aside and added back
# after the HC and Appian updates (sections 9 and 9A.2)
updated_ref_df, ref_blank_key_rows = split_blank_keys(updated_ref_df, KEY_COL)

if len(ref_blank_key_rows):
    print(
        f"Note: {len(ref_blank_key_rows)} Reference row(s) have a blank "
        f"{KEY_COL} - kept as-is, not updated from HC/Appian."
    )


# ============================================================
# 9. UPDATE CHANGED FIELDS FROM HEADCOUNT AND AUDIT
# ============================================================

other_changes = []
blanks_kept = []  # position attributes blank in HC / Appian, Reference value kept
onboarded_count = 0
filled_keys = set()  # vacant last month, filled this month

# Use the existing comments column if present (any case), else create it
comments_col = find_column(updated_ref_df, COMMENTS_RD_COL)
if comments_col is None:
    comments_col = COMMENTS_RD_COL
    updated_ref_df[comments_col] = ""
    print(f"Note: '{COMMENTS_RD_COL}' not found in Reference - column added.")

updated_ref_df = drop_duplicate_keys(updated_ref_df, KEY_COL, "Reference")
hc_unique_df = drop_duplicate_keys(hc_df.dropna(subset=[KEY_COL]), KEY_COL, "HC")

updated_ref_df = updated_ref_df.set_index(KEY_COL)
hc_indexed = hc_unique_df.set_index(KEY_COL)

# --------------------------------------------------------
# VACANT LAST MONTH -> FILLED THIS MONTH (checked per Position Number)
#
# Vacant in PM : Employee Name = Vacant or blank, OR Status = VACANT
# Filled in CM : HC Employee Name is a real name (not Vacant / blank)
#                AND HC Status (if HC has one) is not VACANT
# Every position vacant in PM, and every position whose employee
# changed, is listed on the "Filled-Changed check" sheet with the
# outcome, so each Position Number can be checked.
# --------------------------------------------------------
pm_status_col = find_column_loose(updated_ref_df, STATUS_CANDIDATES)
hc_status_col = find_column_loose(hc_indexed, STATUS_CANDIDATES)
hc_emp_id_col = find_column_loose(hc_indexed, ["Employee ID", "EmployeeID", "Emp ID"])
ref_emp_id_col = find_column_loose(updated_ref_df, ["Employee ID", "EmployeeID", "Emp ID"])
vacant_pm_check = []
vacated_keys = set()  # occupied last month, vacant this month


def is_vacant_or_blank(value):
    return is_vacant(value) or clean_compare_value(value) == ""


common_columns = [
    c for c in updated_ref_df.columns
    if c in hc_indexed.columns and c != new_flag_col  # existing flags are never overwritten
]

for pos in common_keys:
    if pos not in updated_ref_df.index or pos not in hc_indexed.index:
        continue

    old_emp_rd = ""
    if OLD_EMP_RD_COL in updated_ref_df.columns:
        old_emp_rd = clean_compare_value(updated_ref_df.at[pos, OLD_EMP_RD_COL])

    change_record = {
        KEY_COL: clean_compare_value(pos),
        OLD_EMP_RD_COL: old_emp_rd,
    }
    has_change = False

    # Last month's employee / status, read before HC overwrites them
    prev_employee = ""
    if EMPLOYEE_NAME_COL in updated_ref_df.columns:
        prev_employee = updated_ref_df.at[pos, EMPLOYEE_NAME_COL]
    prev_status = (
        clean_compare_value(updated_ref_df.at[pos, pm_status_col])
        if pm_status_col is not None else ""
    )
    vacant_last_month = (
        is_vacant_or_blank(prev_employee) or str(prev_status).upper() == "VACANT"
    )
    prev_emp_id = (
        clean_compare_value(updated_ref_df.at[pos, ref_emp_id_col])
        if ref_emp_id_col is not None else ""
    )

    for col in common_columns:
        old_val = updated_ref_df.at[pos, col]
        new_val = hc_indexed.at[pos, col]

        if values_differ(old_val, new_val):
            # Position attribute blank in HC -> Reference value kept
            if (KEEP_VALUE_WHEN_SOURCE_BLANK and keep_when_blank(col)
                    and clean_compare_value(new_val) == ""):
                blanks_kept.append({
                    KEY_COL: clean_compare_value(pos),
                    EMPLOYEE_NAME_COL: clean_compare_value(
                        hc_indexed.at[pos, EMPLOYEE_NAME_COL]
                        if EMPLOYEE_NAME_COL in hc_indexed.columns else prev_employee
                    ),
                    "Field": col,
                    "Value kept": clean_compare_value(old_val),
                    "Blank in": "HC",
                })
                continue
            change_record[f"{col}_Pre"] = clean_compare_value(old_val)
            change_record[f"{col}_Post"] = clean_compare_value(new_val)
            updated_ref_df.at[pos, col] = clean_compare_value(new_val)
            has_change = True

    # --------------------------------------------------------
    # OLD_EMP_RD
    #
    # If current employee is VACANT: retain previous OLD_EMP_RD.
    # Otherwise: update OLD_EMP_RD with current Employee Name.
    # --------------------------------------------------------
    if (
        OLD_EMP_RD_COL in updated_ref_df.columns
        and EMPLOYEE_NAME_COL in hc_indexed.columns
    ):
        current_employee = clean_compare_value(
            hc_indexed.at[pos, EMPLOYEE_NAME_COL]
        )
        if str(current_employee).strip().lower() != "vacant":
            updated_ref_df.at[pos, OLD_EMP_RD_COL] = current_employee

    # --------------------------------------------------------
    # Comments_RD
    #
    # Vacant last month and occupied this month -> "Onboarded"
    # --------------------------------------------------------
    # A different employee this month (Employee ID changed, or name
    # changed when no ID) is treated the same as vacant -> filled.
    employee_changed = False
    if EMPLOYEE_NAME_COL in hc_indexed.columns and not vacant_last_month:
        cur_name = hc_indexed.at[pos, EMPLOYEE_NAME_COL]
        cur_id = (clean_compare_value(hc_indexed.at[pos, hc_emp_id_col])
                  if hc_emp_id_col is not None else "")
        if not is_vacant_or_blank(cur_name):
            if prev_emp_id != "" and cur_id != "" and not is_vacant(cur_id):
                employee_changed = values_differ(prev_emp_id, cur_id)
            else:
                employee_changed = normalise_text(prev_employee) != normalise_text(cur_name)

    if EMPLOYEE_NAME_COL in hc_indexed.columns and (vacant_last_month or employee_changed):
        current_employee = hc_indexed.at[pos, EMPLOYEE_NAME_COL]
        cur_status = (
            clean_compare_value(hc_indexed.at[pos, hc_status_col])
            if hc_status_col is not None else ""
        )
        filled_now = (
            not is_vacant_or_blank(current_employee)
            and str(cur_status).upper() != "VACANT"
        )
        vacant_pm_check.append({
            KEY_COL: clean_compare_value(pos),
            "Change": "Employee changed" if employee_changed else "Vacant last month",
            "PM Employee ID": prev_emp_id,
            "PM Employee Name": clean_compare_value(prev_employee),
            "PM Status": prev_status,
            "In HC": "YES",
            "HC Employee Name": clean_compare_value(current_employee),
            "HC Employee ID": (clean_compare_value(hc_indexed.at[pos, hc_emp_id_col])
                               if hc_emp_id_col is not None else ""),
            "HC Status": cur_status,
            "Filled this month": "YES" if filled_now else "NO",
        })

        if filled_now:
            change_record[f"{comments_col}_Pre"] = clean_compare_value(
                updated_ref_df.at[pos, comments_col]
            )
            change_record[f"{comments_col}_Post"] = ONBOARDED_COMMENT
            updated_ref_df.at[pos, comments_col] = ONBOARDED_COMMENT
            onboarded_count += 1
            filled_keys.add(pos)
            has_change = True

    # --------------------------------------------------------
    # Occupied last month and vacant this month -> vacated
    # --------------------------------------------------------
    if (
        EMPLOYEE_NAME_COL in hc_indexed.columns
        and not vacant_last_month
        and is_vacant_or_blank(hc_indexed.at[pos, EMPLOYEE_NAME_COL])
    ):
        vacated_keys.add(pos)

    if has_change:
        other_changes.append(change_record)

updated_ref_df = updated_ref_df.reset_index()

other_changes_df = pd.DataFrame(other_changes)

# Positions vacant in PM that are not in this month's HC at all
for _, row in ref_df[ref_df[KEY_COL].isin(removed_keys)].iterrows():
    prev_status = clean_compare_value(row[pm_status_col]) if pm_status_col is not None else ""
    if is_vacant_or_blank(row.get(EMPLOYEE_NAME_COL, "")) or str(prev_status).upper() == "VACANT":
        vacant_pm_check.append({
            KEY_COL: clean_compare_value(row[KEY_COL]),
            "Change": "Vacant last month",
            "PM Employee Name": clean_compare_value(row.get(EMPLOYEE_NAME_COL, "")),
            "PM Status": prev_status,
            "In HC": "NO - position not in HC",
            "Filled this month": "NO",
        })

vacant_pm_check_df = pd.DataFrame(vacant_pm_check, columns=[
    KEY_COL, "Change", "PM Employee ID", "PM Employee Name", "PM Status", "In HC",
    "HC Employee ID", "HC Employee Name", "HC Status", "Filled this month",
])
for _c in ["PM Employee ID", "HC Employee ID"]:
    vacant_pm_check_df[_c] = clean_position_id_series(vacant_pm_check_df[_c])
_vac = vacant_pm_check_df["Change"].eq("Vacant last month")
print(f"Positions vacant last month: {int(_vac.sum())} "
      f"(filled this month: {int((_vac & vacant_pm_check_df['Filled this month'].eq('YES')).sum())})")
print(f"Positions with a different employee this month: {int((~_vac).sum())}")

print(f"Positions marked '{ONBOARDED_COMMENT}' (vacant -> occupied): {onboarded_count}")


# ============================================================
# 9A.0 UPDATE APPIAN POSITION ID FROM NEW_MIGRATIONS
#
# ONLY IF POSITION ID IS BLANK
# ============================================================

new_migrations_df[FRC_COL] = clean_frc_series(new_migrations_df[FRC_COL])
appian_df[FRC_COL] = clean_frc_series(appian_df[FRC_COL])

appian_df[POSITION_ID_COL] = clean_position_id_series(appian_df[POSITION_ID_COL])
new_migrations_df[POSITION_ID_COL] = clean_position_id_series(
    new_migrations_df[POSITION_ID_COL]
)

# Lookup table: where duplicate FRC Codes exist in New_Migrations,
# the first Position ID is used.
migration_lookup = (
    new_migrations_df
    .dropna(subset=[POSITION_ID_COL])
    .drop_duplicates(subset=[FRC_COL], keep="first")
    .set_index(FRC_COL)[POSITION_ID_COL]
)

blank_position = appian_df[POSITION_ID_COL].isna()
before_blank = blank_position.sum()

appian_df.loc[blank_position, POSITION_ID_COL] = (
    appian_df.loc[blank_position, FRC_COL].map(migration_lookup)
)

after_blank = appian_df[POSITION_ID_COL].isna().sum()

print(f"Position ID blanks before update : {before_blank}")
print(f"Position ID blanks after update  : {after_blank}")
print(f"Rows updated from New_Migrations : {before_blank - after_blank}")

matches = appian_df[FRC_COL].isin(new_migrations_df[FRC_COL]).sum()
print(f"Matching FRC Codes found: {matches}")


# ============================================================
# 9A.1 APPIAN DERIVED LOOKUP KEY
#
# If Position ID is blank AND Role Type contains Replacement,
# use Existing Position ID
# ============================================================

APPIAN_KEY_COL = "Position ID"
APPIAN_EXISTING_POS_COL = "Existing Position ID"
ROLE_TYPE_COL = "Role Type"
APPIAN_LOOKUP_KEY_COL = "APPIAN_LOOKUP_KEY"

appian_df[APPIAN_KEY_COL] = clean_position_id_series(appian_df[APPIAN_KEY_COL])
appian_df[APPIAN_EXISTING_POS_COL] = clean_position_id_series(
    appian_df[APPIAN_EXISTING_POS_COL]
)

role_type_clean = clean_text_series(appian_df[ROLE_TYPE_COL], lower=True)

appian_df[APPIAN_LOOKUP_KEY_COL] = appian_df[APPIAN_KEY_COL]

replacement_mask = (
    appian_df[APPIAN_KEY_COL].isna()
    & role_type_clean.str.contains("replacement", na=False)
)
appian_df.loc[replacement_mask, APPIAN_LOOKUP_KEY_COL] = (
    appian_df.loc[replacement_mask, APPIAN_EXISTING_POS_COL]
)


# ============================================================
# 9A.2 UPDATE REFERENCE FROM APPIAN
#
# A separate lookup table is used so that appian_df keeps every
# row for the FRC reconciliation in section 12A.
# ============================================================

_, appian_fields = get_appian_config(appian_config_df)

appian_lookup_df = drop_duplicate_keys(
    appian_df.dropna(subset=[APPIAN_LOOKUP_KEY_COL]),
    APPIAN_LOOKUP_KEY_COL,
    "Appian",
)

updated_ref_df = updated_ref_df.set_index(KEY_COL)
appian_indexed = appian_lookup_df.set_index(APPIAN_LOOKUP_KEY_COL)

appian_common_keys = (
    set(updated_ref_df.index.dropna()) & set(appian_indexed.index.dropna())
)

print(f"Appian Position IDs matched to Reference: {len(appian_common_keys)}")

appian_update_cols = [
    c for c in appian_fields
    if c in updated_ref_df.columns and c in appian_indexed.columns
]

for pos in appian_common_keys:
    for col in appian_update_cols:
        old_val = updated_ref_df.at[pos, col]
        new_val = appian_indexed.at[pos, col]

        if values_differ(old_val, new_val):
            # Position attribute blank in Appian -> Reference value kept
            if (KEEP_VALUE_WHEN_SOURCE_BLANK and keep_when_blank(col)
                    and clean_compare_value(new_val) == ""):
                blanks_kept.append({
                    KEY_COL: clean_compare_value(pos),
                    EMPLOYEE_NAME_COL: (
                        clean_compare_value(updated_ref_df.at[pos, EMPLOYEE_NAME_COL])
                        if EMPLOYEE_NAME_COL in updated_ref_df.columns else ""
                    ),
                    "Field": col,
                    "Value kept": clean_compare_value(old_val),
                    "Blank in": "Appian",
                })
                continue
            updated_ref_df.at[pos, col] = clean_compare_value(new_val)

updated_ref_df = updated_ref_df.reset_index()

blanks_kept_df = pd.DataFrame(
    blanks_kept, columns=[KEY_COL, EMPLOYEE_NAME_COL, "Field", "Value kept", "Blank in"]
)
print(f"Position attributes blank in HC / Appian, Reference value kept: "
      f"{len(blanks_kept_df)} (see 'Blank values kept' sheet)")

# Add back rows with a blank Position Number
updated_ref_df = pd.concat(
    [updated_ref_df, ref_blank_key_rows], ignore_index=True
)

if DROP_REMOVED_POSITIONS:
    updated_ref_df = updated_ref_df[
        ~updated_ref_df[KEY_COL].isin(removed_keys)
    ].reset_index(drop=True)
    print(f"Removed positions dropped from Reference: {len(removed_keys)}")


# ============================================================
# 9B. DERIVE ROLE MOVEMENT
# ============================================================

ROLE_MOVEMENT_COL = "Role Movement"

if FRC_COL in updated_ref_df.columns and ROLE_MOVEMENT_COL in updated_ref_df.columns:
    frc = clean_frc_series(updated_ref_df[FRC_COL])
    role_movement = updated_ref_df[ROLE_MOVEMENT_COL].fillna("").astype(str).str.strip()

    blank_role = role_movement.eq("")

    updated_ref_df.loc[
        blank_role & frc.str.startswith("NEW/BTO"), ROLE_MOVEMENT_COL
    ] = "NEW BTO ROLE"

    updated_ref_df.loc[
        blank_role & frc.str.startswith("NEW/MIGRATIONS"), ROLE_MOVEMENT_COL
    ] = "NEW MIGRATION ROLE"
else:
    print("Role Movement logic skipped because required columns are missing.")


# ============================================================
# 9C. FRC ENDING OLD -> OPCO STATUS COMPLETED
# ============================================================

OPCO_STATUS_COL = "OPCO Status"

if FRC_COL in updated_ref_df.columns and OPCO_STATUS_COL in updated_ref_df.columns:
    frc_clean = clean_frc_series(updated_ref_df[FRC_COL])
    updated_ref_df.loc[frc_clean.str.endswith("_OLD"), OPCO_STATUS_COL] = "COMPLETED"
else:
    print(
        "FRC OLD to OPCO Status logic skipped "
        "because required columns are missing."
    )


# ============================================================
# 9D. DERIVE STATUS_APPROVAL
# ============================================================

SUB_FUNCTION_HEAD_COL = "Sub Function Head Status"
COO_STATUS_COL = "COO Status"
STATUS_APPROVAL_COL = "STATUS_APPROVAL"

if all(
    c in updated_ref_df.columns
    for c in [SUB_FUNCTION_HEAD_COL, COO_STATUS_COL, OPCO_STATUS_COL, STATUS_APPROVAL_COL]
):
    sub_function = clean_text_series(updated_ref_df[SUB_FUNCTION_HEAD_COL], lower=True)
    coo = clean_text_series(updated_ref_df[COO_STATUS_COL], lower=True)
    opco = clean_text_series(updated_ref_df[OPCO_STATUS_COL], lower=True)

    updated_ref_df[STATUS_APPROVAL_COL] = np.select(
        [
            sub_function.isin(["not started", "in progress"]),
            coo.eq("not started"),
            coo.eq("in progress"),
            coo.eq("completed") & opco.eq("completed"),
            coo.eq("completed") & opco.isin(["", "in progress", "not started"]),
        ],
        [
            "AWAITING SUB FUNCTION HEAD APPROVAL",
            "AWAITING COO-BE-LEAD APPROVAL",
            "AWAITING COO-BE-LEAD APPROVAL",
            "OPCO APPROVED",
            "AWAITING APPROVAL FROM OPCO",
        ],
        default="",
    )
else:
    print("STATUS_APPROVAL logic skipped because required columns are missing.")


# ============================================================
# 9D.1 Comments_RD FROM STATUS_APPROVAL
#
# STATUS_APPROVAL = AWAITING COO-BE-LEAD APPROVAL
#                or AWAITING APPROVAL FROM OPCO
#   -> Comments_RD = AWAITING FRC APPROVAL
# ============================================================

AWAITING_FRC_STATUSES = [
    "AWAITING COO-BE-LEAD APPROVAL",
    "AWAITING APPROVAL FROM OPCO",
]
AWAITING_FRC_COMMENT = "AWAITING FRC APPROVAL"

if STATUS_APPROVAL_COL in updated_ref_df.columns:
    awaiting_frc = clean_text_series(updated_ref_df[STATUS_APPROVAL_COL]).isin(
        AWAITING_FRC_STATUSES
    )
    updated_ref_df.loc[awaiting_frc, comments_col] = AWAITING_FRC_COMMENT
    print(
        f"Positions marked '{AWAITING_FRC_COMMENT}' in {comments_col}: "
        f"{awaiting_frc.sum()}"
    )

    # --------------------------------------------------------
    # STATUS_APPROVAL = OPCO APPROVED and Comments_RD blank
    #   -> Comments_RD = RAISE IJP
    # Existing comments are left as they are.
    # --------------------------------------------------------
    opco_approved = clean_text_series(updated_ref_df[STATUS_APPROVAL_COL]).eq(
        "OPCO APPROVED"
    )
    comments_blank = is_blank_series(updated_ref_df[comments_col])
    raise_ijp = opco_approved & comments_blank

    updated_ref_df.loc[raise_ijp, comments_col] = "RAISE IJP"
    print(f"Positions marked 'RAISE IJP' in {comments_col}: {raise_ijp.sum()}")
else:
    print(f"{comments_col} from STATUS_APPROVAL skipped - column missing.")


# ============================================================
# 9D.2 FILLED POSITIONS -> OFFER / CANDIDATE / JOINER FIELDS BLANKED
#
# Vacant in the PM file AND occupied in the current month (same check
# as "Onboarded"):
#   Committed Offer, Candidate Name, Joiner flag, Joiner Month,
#   Comments_RD and Leaver are made blank, and the Status flag is set
#   to PHYSICAL. A change of employee (different Employee ID / name)
#   counts as filled too.
# Runs after the Comments_RD rules (so they end up blank) and before
# HIRING_FLG is derived from the joiner flag (9E).
# Column names are matched ignoring case, spaces, "_" and "-".
# ============================================================

FILLED_CLEAR_FIELDS = {
    "Committed Offer": COMMITTED_OFFER_CANDIDATES,
    "Candidate Name": CANDIDATE_NAME_CANDIDATES,
    "Joiner flag": ["JOINER FLAG", "JOINER_FLG", "JOINERS FLAG", "JOINERS_FLG", "JOINERS", "JOINER"],
    "Joiner Month": JOINER_MONTH_CANDIDATES,
    "Comments_RD": [comments_col],
    "Leaver": LEAVER_CANDIDATES,
}

filled_cleared = []
filled_mask = updated_ref_df[KEY_COL].isin(filled_keys)

for label, candidates in FILLED_CLEAR_FIELDS.items():
    col = (find_column_prefix(updated_ref_df, candidates, "leaver") if label == "Leaver"
           else find_column_loose(updated_ref_df, candidates))
    if col is None:
        print(f"Note: '{label}' column not found - not cleared for filled positions.")
        continue

    had_value = filled_mask & ~is_blank_series(updated_ref_df[col])
    for idx in updated_ref_df.index[had_value]:
        filled_cleared.append({
            KEY_COL: updated_ref_df.at[idx, KEY_COL],
            EMPLOYEE_NAME_COL: (updated_ref_df.at[idx, EMPLOYEE_NAME_COL]
                                if EMPLOYEE_NAME_COL in updated_ref_df.columns else ""),
            "Field": col,
            "Pre": clean_compare_value(updated_ref_df.at[idx, col]),
            "Post": "",
        })
    updated_ref_df.loc[filled_mask, col] = None
    print(f"'{col}' blanked for filled positions: {int(had_value.sum())} value(s) cleared")

# Status flag -> PHYSICAL (before the Leaver rule in 9F, so a filled
# position is no longer treated as VACANT)
filled_status_col = find_column_loose(updated_ref_df, STATUS_CANDIDATES)
if filled_status_col is None:
    print("Note: status column not found - not set to "
          f"{FILLED_STATUS_VALUE} for filled positions.")
else:
    to_change = filled_mask & clean_text_series(updated_ref_df[filled_status_col]).ne(FILLED_STATUS_VALUE)
    for idx in updated_ref_df.index[to_change]:
        filled_cleared.append({
            KEY_COL: updated_ref_df.at[idx, KEY_COL],
            EMPLOYEE_NAME_COL: (updated_ref_df.at[idx, EMPLOYEE_NAME_COL]
                                if EMPLOYEE_NAME_COL in updated_ref_df.columns else ""),
            "Field": filled_status_col,
            "Pre": clean_compare_value(updated_ref_df.at[idx, filled_status_col]),
            "Post": FILLED_STATUS_VALUE,
        })
    updated_ref_df.loc[filled_mask, filled_status_col] = FILLED_STATUS_VALUE
    print(f"'{filled_status_col}' set to {FILLED_STATUS_VALUE} for filled positions: "
          f"{int(to_change.sum())}")

# Occupied last month, vacant this month -> Status flag VACANT
vacated_mask = updated_ref_df[KEY_COL].isin(vacated_keys)
if filled_status_col is not None:
    to_vacant = vacated_mask & clean_text_series(updated_ref_df[filled_status_col]).ne(VACANT_STATUS_VALUE)
    for idx in updated_ref_df.index[to_vacant]:
        filled_cleared.append({
            KEY_COL: updated_ref_df.at[idx, KEY_COL],
            EMPLOYEE_NAME_COL: (updated_ref_df.at[idx, EMPLOYEE_NAME_COL]
                                if EMPLOYEE_NAME_COL in updated_ref_df.columns else ""),
            "Field": filled_status_col,
            "Pre": clean_compare_value(updated_ref_df.at[idx, filled_status_col]),
            "Post": VACANT_STATUS_VALUE,
        })
    updated_ref_df.loc[vacated_mask, filled_status_col] = VACANT_STATUS_VALUE
    print(f"'{filled_status_col}' set to {VACANT_STATUS_VALUE} for positions vacated "
          f"this month: {int(to_vacant.sum())}")

# Occupied last month, vacant this month -> Leaver cleared
vacated_leaver_col = find_column_prefix(updated_ref_df, LEAVER_CANDIDATES, "leaver")
if vacated_leaver_col is None:
    print("Note: no Leaver column found - not cleared for vacated positions.")
else:
    had_leaver = vacated_mask & ~is_blank_series(updated_ref_df[vacated_leaver_col])
    for idx in updated_ref_df.index[had_leaver]:
        filled_cleared.append({
            KEY_COL: updated_ref_df.at[idx, KEY_COL],
            EMPLOYEE_NAME_COL: (updated_ref_df.at[idx, EMPLOYEE_NAME_COL]
                                if EMPLOYEE_NAME_COL in updated_ref_df.columns else ""),
            "Field": vacated_leaver_col,
            "Pre": clean_compare_value(updated_ref_df.at[idx, vacated_leaver_col]),
            "Post": "",
        })
    updated_ref_df.loc[vacated_mask, vacated_leaver_col] = None
    print(f"'{vacated_leaver_col}' cleared for positions vacated this month: "
          f"{int(had_leaver.sum())}")

# --------------------------------------------------------
# NEW POSITIONS (and existing positions whose Status flag is blank
# in the PM file) -> Status flag PHYSICAL when the HC shows an
# employee, VACANT when it shows no employee (Vacant / Unspecified,
# or Status = VACANT in HC). Positions not in this month's HC are
# left as they are.
# --------------------------------------------------------
_status_reason = {}  # Position Number -> "New position" / "Status blank in PM"
if filled_status_col is None or EMPLOYEE_NAME_COL not in updated_ref_df.columns:
    print("Note: status / Employee Name column not found - status flag not set "
          "for new positions.")
else:
    hc_vacant_status_keys = set()
    if hc_status_col is not None:
        hc_vacant_status_keys = set(
            hc_indexed.index[clean_text_series(hc_indexed[hc_status_col]).eq("VACANT")]
        )
    status_blank = is_blank_series(updated_ref_df[filled_status_col])
    is_new = updated_ref_df[KEY_COL].isin(new_keys)
    no_employee = (
        updated_ref_df[EMPLOYEE_NAME_COL].apply(is_vacant_or_blank)
        | updated_ref_df[KEY_COL].isin(hc_vacant_status_keys)
    )
    to_set = status_blank & updated_ref_df[KEY_COL].isin(hc_keys)
    for idx in updated_ref_df.index[to_set]:
        new_status = VACANT_STATUS_VALUE if no_employee[idx] else FILLED_STATUS_VALUE
        filled_cleared.append({
            KEY_COL: updated_ref_df.at[idx, KEY_COL],
            EMPLOYEE_NAME_COL: updated_ref_df.at[idx, EMPLOYEE_NAME_COL],
            "Field": filled_status_col,
            "Pre": "",
            "Post": new_status,
        })
        updated_ref_df.at[idx, filled_status_col] = new_status
        _status_reason[updated_ref_df.at[idx, KEY_COL]] = (
            "New position" if is_new[idx] else "Status blank in PM"
        )
    print(f"'{filled_status_col}' set for new positions: {int((to_set & is_new).sum())} "
          f"({int((to_set & is_new & ~no_employee).sum())} {FILLED_STATUS_VALUE}, "
          f"{int((to_set & is_new & no_employee).sum())} {VACANT_STATUS_VALUE}); "
          f"set where blank in PM file: {int((to_set & ~is_new).sum())}")

filled_cleared_df = pd.DataFrame(
    filled_cleared, columns=[KEY_COL, EMPLOYEE_NAME_COL, "Field", "Pre", "Post"]
)


# ============================================================
# 9D.3 Comments_RD = OFFERED -> COMMITTED OFFER = YES
#
# Runs after 9D.2 so a position filled this month (Comments_RD just
# blanked) does not get the flag back. All positions with Committed
# Offer = YES are listed on the "Committed Offers" sheet.
# ============================================================

committed_offer_col = find_column_loose(updated_ref_df, COMMITTED_OFFER_CANDIDATES)
committed_offers_df = pd.DataFrame()

if committed_offer_col is None:
    print("Note: 'Committed Offer' column not found - OFFERED rule skipped.")
else:
    offered = clean_text_series(updated_ref_df[comments_col]).eq(OFFERED_COMMENT)
    newly_set = offered & clean_text_series(updated_ref_df[committed_offer_col]).ne("YES")
    updated_ref_df.loc[offered, committed_offer_col] = "YES"
    print(f"'{committed_offer_col}' set to YES where {comments_col} = {OFFERED_COMMENT}: "
          f"{int(offered.sum())} position(s), {int(newly_set.sum())} newly set")

    has_offer = clean_text_series(updated_ref_df[committed_offer_col]).eq("YES")
    offer_cols = [
        (KEY_COL, KEY_COL),
        ("Position Title", find_column_loose(updated_ref_df, ["Position Title", "Job Title"])),
        ("GCB", find_column_loose(updated_ref_df, GCB_CANDIDATES)),
        ("Work Location", find_column_loose(updated_ref_df, WORK_LOCATION_CANDIDATES)),
        ("Functional Manager", find_functional_manager_col(updated_ref_df)),
        (OLD_EMP_RD_COL, find_column_loose(updated_ref_df, [OLD_EMP_RD_COL])),
        ("Candidate Name", find_column_loose(updated_ref_df, CANDIDATE_NAME_CANDIDATES)),
        ("Joiner Month", find_column_loose(updated_ref_df, JOINER_MONTH_CANDIDATES)),
        ("MR_Lead", find_column_loose(updated_ref_df, MR_LEAD_CANDIDATES)),
        (comments_col, comments_col),
    ]
    committed_offers_df = pd.DataFrame({
        label: (updated_ref_df.loc[has_offer, col].values if col is not None else "")
        for label, col in offer_cols
    })
    committed_offers_df["Set this run"] = np.where(newly_set[has_offer].values, "YES", "")
    missing_cols = [label for label, col in offer_cols if col is None]
    if missing_cols:
        print(f"Note: column(s) not found for the Committed Offers sheet: {', '.join(missing_cols)}")
    print(f"Committed Offers sheet: {len(committed_offers_df)} position(s)")


# ============================================================
# 9E. UPDATE HIRING_FLG
#
# YES if:
#   HIRING AGAINST PHYSICAL_RD = YES
#   OR JOINERS = YES
# ============================================================

HIRING_FLG_COL = "HIRING_FLG"
HIRING_AGAINST_COL = "HIRING AGAINST PHYSICAL_RD"
JOINERS_COL = "JOINERS"
JOINER_FLAG_CANDIDATES = ["JOINER FLAG", "JOINER_FLG", "JOINERS FLAG", "JOINERS_FLG", "JOINERS", "JOINER"]

# ------------------------------------------------------------
# Vacant in the PM file AND flag = YES in the PM file AND filled
# in the current month (same check as "Onboarded") -> flag removed.
# HIRING AGAINST PHYSICAL_RD is cleared here, before HIRING_FLG is
# derived from it; HIRING_FLG itself is cleared in 9E.1.
# ------------------------------------------------------------
flags_removed = []
sfh_demise_keys = set()  # filled in 9E: DEMISE positions keep their flags
demise_check_rows = []   # "DEMISE check" sheet (every Hyderabad position)


def remove_flag_for_filled(col):
    pm_col = find_column(ref_df, col)
    cur_col = find_column(updated_ref_df, col)
    if pm_col is None or cur_col is None:
        print(f"{col} removal for filled positions skipped - column missing.")
        return

    pm_yes = set(ref_df.loc[clean_text_series(ref_df[pm_col]).eq("YES"), KEY_COL].dropna())
    mask = updated_ref_df[KEY_COL].isin((filled_keys & pm_yes) - sfh_demise_keys)

    names = (
        updated_ref_df.loc[mask, EMPLOYEE_NAME_COL]
        if EMPLOYEE_NAME_COL in updated_ref_df.columns
        else pd.Series("", index=updated_ref_df.index[mask])
    )
    for pos, name in zip(updated_ref_df.loc[mask, KEY_COL], names):
        flags_removed.append({
            KEY_COL: pos, EMPLOYEE_NAME_COL: name,
            "Field": cur_col, "Pre": "YES", "Post": "",
        })

    updated_ref_df.loc[mask, cur_col] = None
    print(f"{cur_col} removed (vacant last month, filled this month): {int(mask.sum())}")


remove_flag_for_filled(HIRING_AGAINST_COL)

if HIRING_FLG_COL in updated_ref_df.columns:
    blank = pd.Series("", index=updated_ref_df.index)

    hiring_against = (
        clean_text_series(updated_ref_df[HIRING_AGAINST_COL])
        if HIRING_AGAINST_COL in updated_ref_df.columns
        else blank
    )
    joiners = (
        clean_text_series(updated_ref_df[JOINERS_COL])
        if JOINERS_COL in updated_ref_df.columns
        else blank
    )

    if RESET_HIRING_FLG:
        updated_ref_df[HIRING_FLG_COL] = ""

    updated_ref_df.loc[
        hiring_against.eq("YES") | joiners.eq("YES"), HIRING_FLG_COL
    ] = "YES"

    # --------------------------------------------------------
    # Position in this month's HC AND Sub Function Head Status was
    # BLANK in the PM file AND is Completed in Appian now
    # -> HIRING_FLG = YES. Any other PM value (In Progress, Completed,
    # ...) means no change.
    # Appian is matched on its lookup key (Position ID, or Existing
    # Position ID for replacements); if Appian has no such column the
    # Reference's own Sub Function Head Status is used instead.
    # --------------------------------------------------------
    sfh_appian_col = find_column_loose(appian_lookup_df, [SUB_FUNCTION_HEAD_COL, "Sub Function Head"])
    if sfh_appian_col is not None:
        sfh_done_keys = set(
            appian_lookup_df.loc[
                clean_text_series(appian_lookup_df[sfh_appian_col]).eq("COMPLETED"),
                APPIAN_LOOKUP_KEY_COL,
            ].dropna()
        )
        sfh_source = f"Appian '{sfh_appian_col}'"
    elif SUB_FUNCTION_HEAD_COL in updated_ref_df.columns:
        sfh_done_keys = set(
            updated_ref_df.loc[
                clean_text_series(updated_ref_df[SUB_FUNCTION_HEAD_COL]).eq("COMPLETED"), KEY_COL
            ].dropna()
        )
        sfh_source = f"Reference '{SUB_FUNCTION_HEAD_COL}' (Appian column not found)"
    else:
        sfh_done_keys = set()
        sfh_source = None

    # Only rows that were BLANK in the PM file qualify
    pm_sfh_col = find_column_loose(ref_df, [SUB_FUNCTION_HEAD_COL, "Sub Function Head"])
    if pm_sfh_col is not None:
        pm_sfh_blank_keys = set(ref_df.loc[is_blank_series(ref_df[pm_sfh_col]), KEY_COL].dropna())
    else:
        pm_sfh_blank_keys = set(ref_df[KEY_COL].dropna())
        print(f"Note: no '{SUB_FUNCTION_HEAD_COL}' in the PM file - treated as blank for all rows.")

    sfh_set_rows = []
    sfh_set_keys = set()       # every position the rule applied to (both outcomes)
    sfh_std_keys = set()       # hiring outcome (HIRING_FLG / LEAVER / JOINER = YES)
    sfh_demise_keys = set()    # DEMISE outcome (Work Location in DEMISE_LOCATIONS)
    sfh_phys_keys = set()      # hiring outcome with Status = PHYSICAL (HIRING AGAINST = YES)
    sfh_leaver_yes_keys = set()  # LEAVER set to YES (occupied positions only)
    sfh_leaver_col = find_column_prefix(updated_ref_df, LEAVER_CANDIDATES, "leaver")
    sfh_joiner_col = find_column_loose(updated_ref_df, JOINER_FLAG_CANDIDATES)
    sfh_hap_col = find_column_loose(updated_ref_df, [HIRING_AGAINST_COL, "HIRING_AGAINST_PHYSICAL_RD"])
    sfh_status_col = find_column_loose(updated_ref_df, STATUS_CANDIDATES)
    sfh_loc_col = find_column_loose(updated_ref_df, WORK_LOCATION_CANDIDATES)
    pm_hiring_col_ref = find_column(ref_df, HIRING_FLG_COL)
    if sfh_source is None:
        print("Note: no Sub Function Head Status column in Appian or Reference - rule skipped.")
    else:
        sfh_mask = updated_ref_df[KEY_COL].isin(sfh_done_keys & hc_keys & pm_sfh_blank_keys)

        # Location like Hyderabad (DEMISE_LOCATIONS) -> DEMISE outcome.
        # Every location-like column is checked (Work Location, Location,
        # City, Site ...), in the current values and in the PM file.
        loc_cols = location_columns(updated_ref_df)
        in_demise_loc = pd.Series(False, index=updated_ref_df.index)
        for _c in loc_cols:
            in_demise_loc |= updated_ref_df[_c].apply(is_demise_location)
        pm_demise_loc_keys = set()
        for _c in location_columns(ref_df):
            pm_demise_loc_keys |= set(ref_df.loc[ref_df[_c].apply(is_demise_location), KEY_COL].dropna())
        in_demise_loc |= updated_ref_df[KEY_COL].isin(pm_demise_loc_keys)
        if not loc_cols:
            print(f"Note: no location column found in the Reference - {DEMISE_VALUE} outcome "
                  f"({' / '.join(DEMISE_LOCATIONS)}) cannot be applied.")
        else:
            print(f"Location columns checked for {' / '.join(DEMISE_LOCATIONS)}: "
                  f"{', '.join(loc_cols)} - {int(in_demise_loc.sum())} position(s) match")

        # Already DEMISE in the PM file -> left alone
        already_demise_keys = set()
        if pm_hiring_col_ref is not None:
            already_demise_keys = set(
                ref_df.loc[clean_text_series(ref_df[pm_hiring_col_ref]).eq(DEMISE_VALUE), KEY_COL].dropna()
            )
        if DEMISE_ONLY_NEWLY_APPROVED:
            demise_trigger = sfh_mask
        else:
            demise_trigger = updated_ref_df[KEY_COL].isin((sfh_done_keys & hc_keys) - already_demise_keys)
        sfh_demise = demise_trigger & in_demise_loc
        sfh_std = sfh_mask & ~in_demise_loc
        if sfh_hap_col is not None and sfh_status_col is not None:
            sfh_phys = sfh_std & clean_text_series(updated_ref_df[sfh_status_col]).eq(FILLED_STATUS_VALUE)
        else:
            sfh_phys = pd.Series(False, index=updated_ref_df.index)
            print(f"Note: '{HIRING_AGAINST_COL}' or status column not found - not set for "
                  "Sub Function Head positions.")
        sfh_std_keys = set(updated_ref_df.loc[sfh_std, KEY_COL])
        sfh_demise_keys = set(updated_ref_df.loc[sfh_demise, KEY_COL])
        sfh_set_keys = sfh_std_keys | sfh_demise_keys
        sfh_phys_keys = set(updated_ref_df.loc[sfh_phys, KEY_COL])
        hiring_flg_now = clean_text_series(updated_ref_df[HIRING_FLG_COL])
        std_new = int((sfh_std & hiring_flg_now.ne("YES")).sum())
        demise_new = int((sfh_demise & hiring_flg_now.ne(DEMISE_VALUE)).sum())

        # Vacant position (Status = VACANT or no employee): LEAVER stays
        # blank (a vacant position is never a leaver) and, for DEMISE,
        # Comments_RD = DEMISE
        vacant_now = pd.Series(False, index=updated_ref_df.index)
        if sfh_status_col is not None:
            vacant_now |= clean_text_series(updated_ref_df[sfh_status_col]).eq(VACANT_STATUS_VALUE)
        if EMPLOYEE_NAME_COL in updated_ref_df.columns:
            vacant_now |= updated_ref_df[EMPLOYEE_NAME_COL].apply(is_vacant_or_blank)
        sfh_demise_vacant = sfh_demise & vacant_now
        sfh_vacant = (sfh_std | sfh_demise) & vacant_now
        sfh_leaver_yes_keys = set(updated_ref_df.loc[(sfh_std | sfh_demise) & ~vacant_now, KEY_COL])

        # (label, column, new value) per outcome; None = made blank
        std_spec = [
            (HIRING_FLG_COL, HIRING_FLG_COL, "YES"),
            ("LEAVER", sfh_leaver_col, "YES"),
            ("JOINER", sfh_joiner_col, "YES"),
        ]
        demise_spec = [
            (HIRING_FLG_COL, HIRING_FLG_COL, DEMISE_VALUE),
            ("LEAVER", sfh_leaver_col, "YES"),
            (COMMENTS_RD_COL, comments_col, DEMISE_COMMENT),
            (NEW_POSITION_FLAG_COL, new_flag_col, DEMISE_VALUE),
            (HIRING_AGAINST_COL, sfh_hap_col, None),
        ]
        report_cols = [
            (HIRING_FLG_COL, HIRING_FLG_COL), ("LEAVER", sfh_leaver_col), ("JOINER", sfh_joiner_col),
            (HIRING_AGAINST_COL, sfh_hap_col), (COMMENTS_RD_COL, comments_col),
            (NEW_POSITION_FLAG_COL, new_flag_col),
        ]
        for idx in updated_ref_df.index[sfh_std | sfh_demise]:
            is_demise = bool(sfh_demise[idx])
            is_vacant_pos = bool(sfh_vacant[idx])
            spec = list(demise_spec if is_demise else std_spec)
            if is_vacant_pos:
                # Vacant: LEAVER blank; DEMISE comment becomes DEMISE_VACANT_COMMENT
                spec = [(l, c, None if l == "LEAVER"
                         else DEMISE_VACANT_COMMENT if (is_demise and l == COMMENTS_RD_COL)
                         else v) for l, c, v in spec]
            if not is_demise and sfh_phys[idx]:
                spec.append((HIRING_AGAINST_COL, sfh_hap_col, "YES"))
            row = {
                KEY_COL: updated_ref_df.at[idx, KEY_COL],
                EMPLOYEE_NAME_COL: (updated_ref_df.at[idx, EMPLOYEE_NAME_COL]
                                    if EMPLOYEE_NAME_COL in updated_ref_df.columns else ""),
                "Work Location": (clean_compare_value(updated_ref_df.at[idx, sfh_loc_col])
                                  if sfh_loc_col is not None else ""),
                "Outcome": ((DEMISE_VALUE if is_demise else f"{HIRING_FLG_COL} = YES")
                            + (" (vacant)" if is_vacant_pos else "")),
            }
            for label, col in report_cols:
                pre = clean_compare_value(updated_ref_df.at[idx, col]) if col is not None else ""
                row[f"{label}_Pre"] = pre
                row[f"{label}_Post"] = pre  # unchanged unless set below
            for label, col, value in spec:
                if col is None:
                    row[f"{label}_Post"] = "(column not found)"
                    continue
                row[f"{label}_Post"] = "" if value is None else value
                updated_ref_df.at[idx, col] = value
            sfh_set_rows.append(row)

        for label, col, _ in std_spec + demise_spec:
            if col is None:
                print(f"Note: no {label} column found - not updated for Sub Function Head positions.")
        print(f"{HIRING_FLG_COL} / LEAVER / JOINER = YES where Sub Function Head Status was blank "
              f"in PM, Completed now ({sfh_source}) and position in HC: {int(sfh_std.sum())} "
              f"position(s), {std_new} newly set "
              f"(Completed but not blank in PM, ignored: {len((sfh_done_keys & hc_keys) - pm_sfh_blank_keys)})")
        print(f"'{sfh_hap_col}' set to YES for those with '{sfh_status_col}' = "
              f"{FILLED_STATUS_VALUE}: {int(sfh_phys.sum())}")
        print(f"{DEMISE_VALUE} outcome (location like {' / '.join(DEMISE_LOCATIONS)}, "
              f"Sub Function Head Completed{' and blank in PM' if DEMISE_ONLY_NEWLY_APPROVED else ''}, "
              f"in HC, not already {DEMISE_VALUE}): {int(sfh_demise.sum())} position(s), "
              f"{demise_new} newly set - {HIRING_FLG_COL} and {NEW_POSITION_FLAG_COL} = "
              f"{DEMISE_VALUE}, LEAVER = YES (blank if vacant), {COMMENTS_RD_COL} = '{DEMISE_COMMENT}' "
              f"(vacant: '{DEMISE_VACANT_COMMENT}', {int(sfh_demise_vacant.sum())} position(s)), "
              f"{HIRING_AGAINST_COL} blank")
        print(f"LEAVER left blank for vacant Sub Function Head positions "
              f"(a vacant position is never a leaver): {int(sfh_vacant.sum())}")

        # "DEMISE check" sheet: every position whose location is like
        # Hyderabad, with the reason it was or was not demised
        if sfh_appian_col is not None:
            sfh_now = dict(zip(appian_lookup_df[APPIAN_LOOKUP_KEY_COL], appian_lookup_df[sfh_appian_col]))
        elif SUB_FUNCTION_HEAD_COL in updated_ref_df.columns:
            sfh_now = dict(zip(updated_ref_df[KEY_COL], updated_ref_df[SUB_FUNCTION_HEAD_COL]))
        else:
            sfh_now = {}
        pm_sfh = dict(zip(ref_df[KEY_COL], ref_df[pm_sfh_col])) if pm_sfh_col is not None else {}
        pm_hiring_pm = (dict(zip(ref_df[KEY_COL], ref_df[pm_hiring_col_ref]))
                        if pm_hiring_col_ref is not None else {})
        pm_loc = {}
        for _c in location_columns(ref_df):
            for _k, _v in zip(ref_df[KEY_COL], ref_df[_c]):
                if pd.notna(_k) and clean_compare_value(_v) != "":
                    pm_loc[_k] = (pm_loc[_k] + " | " if _k in pm_loc else "") + str(clean_compare_value(_v))
        for idx in updated_ref_df.index[in_demise_loc]:
            k = updated_ref_df.at[idx, KEY_COL]
            if pd.isna(k):
                continue
            if k in sfh_demise_keys:
                outcome = (f"{DEMISE_VALUE} applied - vacant, {COMMENTS_RD_COL} = {DEMISE_VACANT_COMMENT}, "
                           "LEAVER blank"
                           if sfh_demise_vacant[idx] else
                           f"{DEMISE_VALUE} applied - {COMMENTS_RD_COL} = {DEMISE_COMMENT}")
            elif k not in hc_keys:
                outcome = "Not applied - position not in this month's HC"
            elif k not in sfh_done_keys:
                outcome = f"Not applied - Sub Function Head Status not Completed ({sfh_source})"
            elif k in already_demise_keys:
                outcome = f"Not applied - already {DEMISE_VALUE} in PM file"
            elif DEMISE_ONLY_NEWLY_APPROVED and k not in pm_sfh_blank_keys:
                outcome = ("Not applied - Sub Function Head Status not blank in PM file "
                           "(DEMISE_ONLY_NEWLY_APPROVED = True)")
            else:
                outcome = "Not applied"
            demise_check_rows.append({
                KEY_COL: k,
                EMPLOYEE_NAME_COL: (updated_ref_df.at[idx, EMPLOYEE_NAME_COL]
                                    if EMPLOYEE_NAME_COL in updated_ref_df.columns else ""),
                "Location": " | ".join(
                    str(clean_compare_value(updated_ref_df.at[idx, c])) for c in loc_cols
                    if clean_compare_value(updated_ref_df.at[idx, c]) != ""
                ),
                "Location (PM)": pm_loc.get(k, ""),
                "In HC": "YES" if k in hc_keys else "NO",
                "Sub Function Head Status (PM)": clean_compare_value(pm_sfh.get(k, "")),
                "Sub Function Head Status (now)": clean_compare_value(sfh_now.get(k, "")),
                f"{HIRING_FLG_COL} (PM)": clean_compare_value(pm_hiring_pm.get(k, "")),
                "Outcome": outcome,
            })
        for _o, _n in pd.Series([r["Outcome"] for r in demise_check_rows]).value_counts().items():
            print(f"  {DEMISE_VALUE} check - {_o}: {_n}")
else:
    sfh_set_rows = []
    sfh_set_keys = set()
    sfh_std_keys = set()
    sfh_demise_keys = set()
    sfh_phys_keys = set()
    sfh_leaver_yes_keys = set()
    sfh_leaver_col = sfh_joiner_col = sfh_hap_col = None
    print(f"{HIRING_FLG_COL} not found. Skipping HIRING_FLG update.")

hiring_flag_sfh_df = pd.DataFrame(sfh_set_rows, columns=[
    KEY_COL, EMPLOYEE_NAME_COL, "Work Location", "Outcome",
    f"{HIRING_FLG_COL}_Pre", f"{HIRING_FLG_COL}_Post",
    "LEAVER_Pre", "LEAVER_Post", "JOINER_Pre", "JOINER_Post",
    f"{HIRING_AGAINST_COL}_Pre", f"{HIRING_AGAINST_COL}_Post",
    f"{COMMENTS_RD_COL}_Pre", f"{COMMENTS_RD_COL}_Post",
    f"{NEW_POSITION_FLAG_COL}_Pre", f"{NEW_POSITION_FLAG_COL}_Post",
])
demise_check_df = pd.DataFrame(demise_check_rows, columns=[
    KEY_COL, EMPLOYEE_NAME_COL, "Location", "Location (PM)", "In HC", "Sub Function Head Status (PM)",
    "Sub Function Head Status (now)", f"{HIRING_FLG_COL} (PM)", "Outcome",
])


# ============================================================
# 9E.1 FILLED POSITIONS -> HIRING_FLG REMOVED
#
# Same rule as HIRING AGAINST PHYSICAL_RD above. Runs after 9E so
# the flag is not set back to YES.
# ============================================================

remove_flag_for_filled(HIRING_FLG_COL)

hiring_flag_removed_df = pd.DataFrame(
    flags_removed, columns=[KEY_COL, EMPLOYEE_NAME_COL, "Field", "Pre", "Post"]
)


# ------------------------------------------------------------
# One row per position: Field / Pre / Post rows become
# <Field>_Pre / <Field>_Post columns, plus a "Change" column
# (Vacant -> Filled, Employee changed, Vacated).
# ------------------------------------------------------------
_change_reason = dict(_status_reason)  # New position / Status blank in PM
for _, r in vacant_pm_check_df.iterrows():
    if r["Filled this month"] == "YES":
        _change_reason[r[KEY_COL]] = (
            "Employee changed" if r["Change"] == "Employee changed" else "Vacant -> Filled"
        )
for k in vacated_keys:
    _change_reason[k] = "Vacated"


def one_row_per_position(df):
    base_cols = [KEY_COL, EMPLOYEE_NAME_COL, "Change"]
    if df.empty:
        return pd.DataFrame(columns=base_cols)
    fields = list(dict.fromkeys(df["Field"]))  # keep first-seen order
    rows = {}
    for _, r in df.iterrows():
        key = clean_compare_value(r[KEY_COL])
        row = rows.setdefault(key, {
            KEY_COL: r[KEY_COL],
            EMPLOYEE_NAME_COL: r[EMPLOYEE_NAME_COL],
            "Change": _change_reason.get(r[KEY_COL], ""),
        })
        row[f"{r['Field']}_Pre"] = r["Pre"]
        row[f"{r['Field']}_Post"] = r["Post"]
    cols = base_cols + [f"{f}_{side}" for f in fields for side in ("Pre", "Post")]
    return pd.DataFrame(list(rows.values())).reindex(columns=cols)


filled_cleared_df = one_row_per_position(filled_cleared_df)
hiring_flag_removed_df = one_row_per_position(hiring_flag_removed_df)


# ============================================================
# 9F. VACANT POSITION -> LEAVER CLEARED (no exceptions)
#
# A vacant position can never be a leaver. Vacant = Status = VACANT,
# or no employee (Employee Name Vacant / Unspecified / blank). Runs
# after every other rule, so LEAVER = YES never stays on a vacant
# position. Column names are matched ignoring case.
# ============================================================

status_col = find_column_loose(updated_ref_df, STATUS_CANDIDATES)
leaver_col = find_column_prefix(updated_ref_df, LEAVER_CANDIDATES, "leaver")

if leaver_col is None:
    print("Leaver clearing skipped - no Leaver column found.")
else:
    vacant_status = pd.Series(False, index=updated_ref_df.index)
    if status_col is not None:
        vacant_status |= clean_text_series(updated_ref_df[status_col]).eq(VACANT_STATUS_VALUE)
    if EMPLOYEE_NAME_COL in updated_ref_df.columns:
        vacant_status |= updated_ref_df[EMPLOYEE_NAME_COL].apply(is_vacant_or_blank)
    had_leaver = vacant_status & ~is_blank_series(updated_ref_df[leaver_col])
    updated_ref_df.loc[vacant_status, leaver_col] = None
    print(f"'{leaver_col}' cleared for vacant positions ('{status_col}' = VACANT or no employee): "
          f"{int(vacant_status.sum())} position(s), {int(had_leaver.sum())} had a value")


# ============================================================
# 9G. LWD FROM MR_LEAVERS_CM (Termination Date)
#
# The CURRENT MONTH Employee ID (after the HC update) is compared
# with Employee ID in the leavers file; on a match, LWD is set to the
# leaver's Termination Date (latest date if listed more than once).
#
# Employees who have already exited (position Vacant / Employee ID
# blank in the current month) are not matched, so no Termination
# Date is added for them.
#
# Only records with Employee Class = "Employee" are compared (taken
# from the leavers file; if it has no Employee Class column, the
# Reference Data's Employee Class is used instead).
# ============================================================

LWD_COL = "LWD"
TERMINATION_DATE_CANDIDATES = [
    "Termination Date", "Terminate Date", "Terminated Date", "Date of Termination",
]
EMPLOYEE_ID_CANDIDATES = ["Employee ID", "EmployeeID", "Emp ID", "Employee Number"]
EMPLOYEE_CLASS_COL = "Employee Class"
EMPLOYEE_CLASS_VALUE = "EMPLOYEE"

lwd_updates_df = pd.DataFrame(columns=[
    KEY_COL, "Employee ID", EMPLOYEE_NAME_COL, "Global Career Band", "Work Location",
    "Functional Manager", f"{LWD_COL}_Pre", f"{LWD_COL}_Post",
])
leavers_unmatched_df = pd.DataFrame()


def first_column(df, candidates):
    for name in candidates:
        col = find_column(df, name)
        if col is not None:
            return col
    return None


if not os.path.exists(leavers_file):
    print(f"LWD update skipped - leavers file not found: {leavers_file}")
else:
    leavers_df = clean_columns(pd.read_excel(leavers_file))

    term_col = first_column(leavers_df, TERMINATION_DATE_CANDIDATES)
    lv_emp_col = first_column(leavers_df, EMPLOYEE_ID_CANDIDATES)
    cur_emp_col = first_column(updated_ref_df, EMPLOYEE_ID_CANDIDATES)

    if term_col is None:
        print("LWD update skipped - no 'Termination Date' column in the leavers file.")
    elif lv_emp_col is None:
        print("LWD update skipped - no 'Employee ID' column in the leavers file.")
    elif cur_emp_col is None:
        print("LWD update skipped - no 'Employee ID' column in the Reference Data Hierarchy.")
    else:
        lwd_col = find_column(updated_ref_df, LWD_COL)
        if lwd_col is None:
            lwd_col = LWD_COL
            updated_ref_df[lwd_col] = None
            print(f"Note: '{LWD_COL}' not found in Reference - column added.")

        # Only Employee Class = Employee
        lv_class_col = find_column(leavers_df, EMPLOYEE_CLASS_COL)
        ref_class_col = find_column(updated_ref_df, EMPLOYEE_CLASS_COL)
        if lv_class_col is not None:
            is_employee = clean_text_series(leavers_df[lv_class_col]).eq(EMPLOYEE_CLASS_VALUE)
            print(f"Leavers with Employee Class = Employee: {int(is_employee.sum())} "
                  f"of {len(leavers_df)} (others ignored for LWD)")
            leavers_df = leavers_df[is_employee].copy()
        elif ref_class_col is None:
            print(f"Note: no '{EMPLOYEE_CLASS_COL}' column in the leavers file or the "
                  "Reference - all leavers compared.")

        # Leaver Employee ID -> latest Termination Date
        leavers_df["_EMP"] = clean_position_id_series(leavers_df[lv_emp_col])
        leavers_df["_TERM_DATE"] = pd.to_datetime(leavers_df[term_col], errors="coerce")
        leaver_dates = (
            leavers_df.dropna(subset=["_EMP", "_TERM_DATE"])
            .sort_values("_TERM_DATE")
            .drop_duplicates("_EMP", keep="last")
            .set_index("_EMP")["_TERM_DATE"]
        )

        # Current month Employee ID of each Reference row. Blank if Vacant,
        # and ignored for positions not in this month's HC (removed positions
        # still carry last month's Employee ID).
        ref_emp = clean_position_id_series(updated_ref_df[cur_emp_col])
        ref_emp[~updated_ref_df[KEY_COL].isin(hc_keys)] = np.nan
        if lv_class_col is None and ref_class_col is not None:
            ref_is_employee = clean_text_series(updated_ref_df[ref_class_col]).eq(EMPLOYEE_CLASS_VALUE)
            ref_emp[~ref_is_employee] = np.nan
            print(f"Reference rows with Employee Class = Employee: {int(ref_is_employee.sum())}")

        hit = ref_emp.isin(leaver_dates.index)
        new_lwd = ref_emp[hit].map(leaver_dates)
        old_lwd = updated_ref_df.loc[hit, lwd_col]
        # boolean array (not a list) so an empty result still keeps the columns
        changed = np.array([values_differ(o, n) for o, n in zip(old_lwd, new_lwd)], dtype=bool)

        # Functional Manager name column (e.g. "Functional Manager Job Level
        # Employee Name"); ID / position-number columns are skipped
        fm_col = next(
            (c for c in updated_ref_df.columns
             if str(c).lower().startswith("functional manager") and "name" in str(c).lower()),
            next((c for c in updated_ref_df.columns
                  if str(c).lower().startswith("functional manager")
                  and not re.search(r"\bid\b|number", str(c).lower())), None),
        )
        hit_rows = updated_ref_df.loc[hit].loc[changed]
        gcb_col = find_column_loose(updated_ref_df, GCB_CANDIDATES)
        loc_col = find_column_loose(updated_ref_df, WORK_LOCATION_CANDIDATES)

        lwd_updates_df = pd.DataFrame({
            KEY_COL: hit_rows[KEY_COL].values,
            "Employee ID": ref_emp[hit][changed].values,
            EMPLOYEE_NAME_COL: (hit_rows[EMPLOYEE_NAME_COL].values
                                if EMPLOYEE_NAME_COL in hit_rows.columns else ""),
            "Global Career Band": hit_rows[gcb_col].values if gcb_col is not None else "",
            "Work Location": hit_rows[loc_col].values if loc_col is not None else "",
            "Functional Manager": hit_rows[fm_col].values if fm_col is not None else "",
            # blank, 0 or placeholder dates -> blank LWD_Pre
            f"{LWD_COL}_Pre": [to_date(v) for v in old_lwd[changed]],
            f"{LWD_COL}_Post": new_lwd[changed].values,
        })
        if fm_col is None:
            print("Note: no Functional Manager name column found for the LWD updates sheet.")
        updated_ref_df.loc[hit, lwd_col] = new_lwd

        not_matched = (
            ~leavers_df["_EMP"].isin(set(ref_emp.dropna()))
            | leavers_df["_TERM_DATE"].isna()
        )
        leavers_unmatched_df = leavers_df[not_matched].drop(columns=["_EMP", "_TERM_DATE"])

        print(f"Leavers in file                  : {len(leavers_df)}")
        print(f"Reference rows matched on Emp ID : {int(hit.sum())}")
        print(f"LWD values changed               : {len(lwd_updates_df)}")
        if len(leavers_unmatched_df):
            print(f"Note: {len(leavers_unmatched_df)} leaver row(s) not in the current month "
                  "(already exited) or without a Termination Date - see 'Leavers not matched' sheet.")


# ============================================================
# 9H. GCB CHECK
#
# Where the blank GCBs in the output come from: already blank in the
# PM file, new positions with no GCB in HC, or GCB blank in HC.
# ============================================================

gcb_col = find_column_loose(updated_ref_df, GCB_CANDIDATES)
if gcb_col is None:
    print("Note: no GCB column found in the Reference.")
else:
    gcb_blank = is_blank_series(updated_ref_df[gcb_col])
    pm_gcb_col = find_column_loose(ref_df, GCB_CANDIDATES)
    pm_gcb_blank_keys = (
        set(ref_df.loc[is_blank_series(ref_df[pm_gcb_col]), KEY_COL].dropna())
        if pm_gcb_col is not None else set()
    )
    hc_gcb_col = find_column_loose(hc_df, GCB_CANDIDATES)
    hc_gcb_blank_keys = (
        set(hc_df.loc[is_blank_series(hc_df[hc_gcb_col]), KEY_COL].dropna())
        if hc_gcb_col is not None else set(hc_df[KEY_COL].dropna())
    )
    _keys = updated_ref_df[KEY_COL]
    print(f"'{gcb_col}' blank in the output: {int(gcb_blank.sum())} position(s) "
          f"(already blank in PM file: {int((gcb_blank & _keys.isin(pm_gcb_blank_keys)).sum())}, "
          f"new positions: {int((gcb_blank & _keys.isin(new_keys)).sum())}, "
          f"blank in HC: {int((gcb_blank & _keys.isin(hc_gcb_blank_keys)).sum())})")
    if hc_gcb_col is None:
        print("Note: no GCB column in the HC file "
              f"({' / '.join(GCB_CANDIDATES)}) - new positions get a blank GCB.")

# ============================================================
# 10. COLUMN SEQUENCING FROM CONFIG
# ============================================================

final_column_order = get_export_column_order(config_df)
updated_ref_df = apply_column_sequence(updated_ref_df, final_column_order)

positions_removed_df = positions_removed_df.reindex(
    columns=updated_ref_df.columns, fill_value=""
)
# "new positions added in CM" shows the new rows as they are in the
# Reference output (status flag, GCB, flags and comments included)
new_positions_df = updated_ref_df[updated_ref_df[KEY_COL].isin(new_keys)].copy()

if not other_changes_df.empty:
    ordered_cols = [
        c for c in [KEY_COL, OLD_EMP_RD_COL] if c in other_changes_df.columns
    ]
    remaining_cols = [c for c in other_changes_df.columns if c not in ordered_cols]
    other_changes_df = other_changes_df[ordered_cols + remaining_cols]
else:
    other_changes_df = pd.DataFrame(columns=[KEY_COL, OLD_EMP_RD_COL])


# ============================================================
# 11. JOB SUMMARY EXCEPTIONS
# ============================================================

job_summary_exclusions = get_job_summary_exclusions(config_df)
expected_job_value = normalise_text(EXPECTED_JOB_SUMMARY)

if JOB_SUMMARY_COL in updated_ref_df.columns:
    job_summary_clean = updated_ref_df[JOB_SUMMARY_COL].apply(normalise_text)
    job_summary_exceptions_df = updated_ref_df[
        (job_summary_clean != expected_job_value)
        & (~job_summary_clean.isin(job_summary_exclusions))
    ].copy()
else:
    job_summary_exceptions_df = pd.DataFrame()


# ============================================================
# 12. DATA TYPE HANDLING
# ============================================================

updated_ref_df = apply_data_types(updated_ref_df)
positions_removed_df = apply_data_types(positions_removed_df)
new_positions_df = apply_data_types(new_positions_df)
other_changes_df = apply_data_types(other_changes_df)
job_summary_exceptions_df = apply_data_types(job_summary_exceptions_df)


# ============================================================
# 12A. FRC RECONCILIATION REPORT
#
# APPIAN FRC vs REFERENCE FRC
# Uses cleaned copies so the exported Reference FRC is unchanged.
# Uses every Appian row, including those without a Position ID.
# ============================================================

if FRC_COL in updated_ref_df.columns and FRC_COL in appian_df.columns:
    ref_frc = clean_frc_series(updated_ref_df[FRC_COL])
    appian_frc = clean_frc_series(appian_df[FRC_COL])

    appian_frc_set = set(appian_frc) - {""}
    ref_frc_set = set(ref_frc) - {""}

    appian_not_in_ref = appian_df[
        appian_frc.isin(appian_frc_set - ref_frc_set)
    ].copy()

    ref_not_in_appian = updated_ref_df[
        ref_frc.isin(ref_frc_set - appian_frc_set)
    ].copy()

    control_summary_df = pd.DataFrame({
        "Metric": [
            "Total FRC in Appian",
            "Total FRC in Reference",
            "Matching FRC",
            "Appian FRC Not in Reference",
            "Reference FRC Not in Appian",
        ],
        "Count": [
            len(appian_frc_set),
            len(ref_frc_set),
            len(appian_frc_set & ref_frc_set),
            len(appian_frc_set - ref_frc_set),
            len(ref_frc_set - appian_frc_set),
        ],
    })

    with pd.ExcelWriter(frc_reconciliation_file, engine="openpyxl") as writer:
        control_summary_df.to_excel(writer, index=False, sheet_name="Control Summary")
        appian_not_in_ref.to_excel(writer, index=False, sheet_name="Appian_FRC_Not_in_Ref")
        ref_not_in_appian.to_excel(writer, index=False, sheet_name="Ref_FRC_Not_in_Appian")

    format_excel(frc_reconciliation_file)
else:
    print("FRC reconciliation skipped - FRC Code column missing.")


# ============================================================
# 12B. CREATE MR HIRING UPDATE FILE
# ============================================================

hiring_update_cols = get_hiring_update_columns(config_df)

if HIRING_FLG_COL in updated_ref_df.columns:
    # If no specific config list found, use all reference columns.
    if hiring_update_cols:
        valid_hiring_cols = [c for c in hiring_update_cols if c in updated_ref_df.columns]
    else:
        valid_hiring_cols = list(updated_ref_df.columns)

    if HIRING_FLG_COL not in valid_hiring_cols:
        valid_hiring_cols.append(HIRING_FLG_COL)

    hiring_update_df = updated_ref_df[
        clean_text_series(updated_ref_df[HIRING_FLG_COL]).eq("YES")
    ][valid_hiring_cols].copy()

    hiring_update_df = apply_data_types(hiring_update_df)

    with pd.ExcelWriter(hiring_update_file, engine="openpyxl") as writer:
        hiring_update_df.to_excel(writer, index=False, sheet_name="MR_Hiring_Update")

    format_excel(hiring_update_file)
else:
    print("MR Hiring Update file skipped.")


# ============================================================
# 13. EXPORT OUTPUT FILES
#
# Config goes only to Reference output.
# New_Migrations is also retained.
# ============================================================

with pd.ExcelWriter(output_reference_file, engine="openpyxl") as writer:
    updated_ref_df.to_excel(writer, index=False, sheet_name="Reference_Data_CM")
    config_df.to_excel(writer, index=False, sheet_name="Config")
    appian_config_df.to_excel(writer, index=False, sheet_name="APPIAN_CONFIG")
    new_migrations_df.to_excel(writer, index=False, sheet_name="New_Migrations")

# ------------------------------------------------------------
# EXCEPTION REPORT
# ------------------------------------------------------------

with pd.ExcelWriter(exception_file, engine="openpyxl") as writer:
    positions_removed_df.to_excel(writer, index=False, sheet_name="positions removed")
    new_positions_df.to_excel(writer, index=False, sheet_name="new positions added in CM")
    other_changes_df.to_excel(writer, index=False, sheet_name="Other changes")
    job_summary_exceptions_df.to_excel(writer, index=False, sheet_name="Job Summary Exceptions")
    apply_data_types(lwd_updates_df).to_excel(writer, index=False, sheet_name="LWD updates")
    apply_data_types(leavers_unmatched_df).to_excel(writer, index=False, sheet_name="Leavers not matched")
    apply_data_types(hiring_flag_removed_df).to_excel(writer, index=False, sheet_name="Hiring flags removed")
    apply_data_types(hiring_flag_sfh_df).to_excel(writer, index=False, sheet_name="Hiring flag set (SFH)")
    apply_data_types(demise_check_df).to_excel(writer, index=False, sheet_name="DEMISE check")
    filled_cleared_df.to_excel(writer, index=False, sheet_name="Filled-Vacated changes")
    vacant_pm_check_df.to_excel(writer, index=False, sheet_name="Filled-Changed check")
    apply_data_types(committed_offers_df).to_excel(writer, index=False, sheet_name="Committed Offers")
    apply_data_types(blanks_kept_df).to_excel(writer, index=False, sheet_name="Blank values kept")


# ============================================================
# 14. FORMAT OUTPUT FILES
# ============================================================

format_excel(output_reference_file)
format_excel(exception_file)

# HIRING_FLG cells that changed this run (vs the PM file) -> light green
hiring_flag_changed_keys = set()
pm_hiring_col = find_column(ref_df, HIRING_FLG_COL)
if pm_hiring_col is not None and HIRING_FLG_COL in updated_ref_df.columns:
    pm_hiring = {
        str(clean_compare_value(k)): str(clean_compare_value(v)).upper()
        for k, v in zip(ref_df[KEY_COL], ref_df[pm_hiring_col]) if pd.notna(k)
    }
    for k, v in zip(updated_ref_df[KEY_COL], updated_ref_df[HIRING_FLG_COL]):
        if pd.isna(k):
            continue
        key = str(clean_compare_value(k))
        if str(clean_compare_value(v)).upper() != pm_hiring.get(key, ""):
            hiring_flag_changed_keys.add(key)

    n_ref = highlight_cells(output_reference_file, "Reference_Data_CM", HIRING_FLG_COL, hiring_flag_changed_keys)
    n_hu = highlight_cells(hiring_update_file, "MR_Hiring_Update", HIRING_FLG_COL, hiring_flag_changed_keys)
    print(f"{HIRING_FLG_COL} changed this run: {len(hiring_flag_changed_keys)} position(s) "
          f"highlighted light green (Reference: {n_ref}, MR Hiring Update: {n_hu})")

# LEAVER / JOINER / HIRING AGAINST PHYSICAL_RD / Comments_RD / NEW_POSITION_FLAG
# set by the Sub Function Head rule (both outcomes) -> light green too
for _col, _ks in ((sfh_leaver_col, sfh_leaver_yes_keys), (sfh_joiner_col, sfh_std_keys),
                  (sfh_hap_col, sfh_phys_keys | sfh_demise_keys),
                  (comments_col, sfh_demise_keys), (new_flag_col, sfh_demise_keys)):
    if _col is not None and _ks:
        _keys = {str(clean_compare_value(k)) for k in _ks}
        highlight_cells(output_reference_file, "Reference_Data_CM", _col, _keys)
        highlight_cells(hiring_update_file, "MR_Hiring_Update", _col, _keys)


# ============================================================
# 15. COMPLETION MESSAGE
# ============================================================

print("")
print("Completed successfully.")
print(f"Reference output created: {output_reference_file}")
print(f"Exception report created: {exception_file}")

if os.path.exists(hiring_update_file):
    print(f"MR Hiring Update file created: {hiring_update_file}")

if os.path.exists(frc_reconciliation_file):
    print(f"FRC Reconciliation report created: {frc_reconciliation_file}")

print(f"Run log saved: {run_log_file}")

# Stop writing to the log file
sys.stdout = original_stream(sys.stdout)
sys.stderr = original_stream(sys.stderr)
